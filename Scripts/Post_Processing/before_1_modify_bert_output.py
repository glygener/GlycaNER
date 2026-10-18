#!/usr/bin/env python3
"""
Apply post-processing rules to BERT_Output rows in an Excel workbook.

Input columns expected:
  pmid, sent_index, sent_text, BERT_Output, start_offset, end_offset

Output:
  A new workbook with BERT_Output/start_offset/end_offset updated.
"""

from __future__ import annotations
from glycaner_paths import STAGE3_INPUT, resource, work  # noqa: E402

import copy
import re
from collections import Counter, OrderedDict
from pathlib import Path
from typing import Any

import openpyxl


INPUT_FILE = Path(
    STAGE3_INPUT
)
OUTPUT_FILE = Path(
    work("bert_output_full_updated.xlsx")
)

REQUIRED_COLUMNS = [
    "pmid",
    "sent_index",
    "sent_text",
    "BERT_Output",
    "start_offset",
    "end_offset",
]

REMOVED_OUTPUT_COLUMNS = {"Modified", "New_detected", "Abbreviated"}
TRACKED_CHANGE_COLUMNS = ("BERT_Output", "start_offset", "end_offset")
PRINT_CHANGED_ROWS = True
MAX_CHANGE_LINES: int | None = None  # Use an integer such as 200 to limit long reports.

CA_125_RE = re.compile(r"\bCA\s+125\b", re.IGNORECASE)
DIGIT_ONLY_RE = re.compile(r"^\s*\d+(?:['\u2032])?\s*$")
O_TION_AFTER_DIGIT_RE = re.compile(
    r"\s*-\s*O\s*-\s*[A-Za-z][A-Za-z-]*tion\b",
    re.IGNORECASE,
)
ALPHA_BETA_GAP_RE = re.compile(r"\s*(?:alpha|beta|\u03b1|\u03b2)\s*", re.IGNORECASE)

LEWIS_PATTERNS = [
    r"(?i)\b\(?((?:S|Sai|SD)-?)?\)?[-]?\(?((?:S|Sai|SD)-?)?\)?[-]?Le[\s^-]?\(?([ABXY])\)?(?!\w)",
    r"(?i)\b\(?((?:S|Sai|SD)-?)?\)?[-]?\(?((?:S|Sai|SD)-?)?\)?[-]?Lewis[\s^-]?\(?([ABXY])\)?(?!\w)",
    r"(?i)\b\(?((?:S|Sai|SD)-?)?\)?[-]?\(?((?:S|Sai|SD)-?)?\)?[-]?Le[\s^-]?\(?([ABXY])\)?-([ABXY])?(?!\w)",
    r"(?i)\b\(?((?:S|Sai|SD)-?)?\)?[-]?\(?((?:S|Sai|SD)-?)?\)?[-]?Lewis[\s^-]?\(?([ABXY])\)?-([ABXY])?(?!\w)",
    r"(?i)\b\(?((?:S|Sai|SD)-?)?\)?[-]?\(?((?:S|Sai|SD)-?)?\)?[-]?Lewis[\s^-]?\(?([ABXY])\)?-antigen(s)?(?!\w)",
    r"(?i)\b(?:S|Sai|SD)?(lewis|le)\s?\(?([ABXY])-?(?:IV|I{1,3})?\)?(?!\w)",
    r"(?i)\bA(Lewis|Le)[\s^()-]*[ABXYD]\b",
    r"\bSLX\b",
]
LEWIS_REGEXES = [re.compile(pattern) for pattern in LEWIS_PATTERNS]

ADDITIONAL_EXACT_TERMS = [
    "CD10",
    "CD26",
    "CD29",
    "CD44",
    "CD73",
    "CD90",
    "CD105",
    "CD166",
]
ADDITIONAL_TERM_PATTERNS = [
    rf"(?<![A-Za-z0-9]){re.escape(term)}(?![A-Za-z0-9])"
    for term in ADDITIONAL_EXACT_TERMS
] + [
    r"(?<![A-Za-z0-9])(?:Sulfatide|sulfatide)(?![A-Za-z0-9])",
]
ADDITIONAL_TERM_REGEXES = [
    re.compile(pattern) for pattern in ADDITIONAL_TERM_PATTERNS
]


def is_blank(value: Any) -> bool:
    return value is None or str(value).strip() == ""


def compact_value(value: Any, max_length: int = 180) -> str:
    if value is None:
        return "None"
    text = str(value).replace("\n", "\\n")
    if len(text) > max_length:
        text = text[: max_length - 3] + "..."
    return repr(text)


def tracked_values(row: dict[str, Any]) -> dict[str, Any]:
    return {column: row.get(column) for column in TRACKED_CHANGE_COLUMNS}


def audit_snapshot(row: dict[str, Any]) -> dict[str, Any]:
    return {
        "input_excel_row": row.get("_excel_row"),
        "output_excel_row": row.get("_output_excel_row"),
        "pmid": row.get("pmid"),
        "sent_index": row.get("sent_index"),
        "sent_text": row.get("sent_text"),
        "tracked": tracked_values(row),
    }


def tracked_values_changed(before: dict[str, Any], after: dict[str, Any]) -> bool:
    before_values = before["tracked"]
    after_values = after["tracked"]
    return any(
        before_values.get(column) != after_values.get(column)
        for column in TRACKED_CHANGE_COLUMNS
    )


def describe_field_changes(before: dict[str, Any], after: dict[str, Any]) -> str:
    changes = []
    before_values = before["tracked"]
    after_values = after["tracked"]

    for column in TRACKED_CHANGE_COLUMNS:
        before_value = before_values.get(column)
        after_value = after_values.get(column)
        if before_value != after_value:
            changes.append(
                f"{column}: {compact_value(before_value)} -> {compact_value(after_value)}"
            )

    return "; ".join(changes)


def describe_tracked_values(snapshot: dict[str, Any]) -> str:
    values = snapshot["tracked"]
    return (
        f"BERT_Output={compact_value(values.get('BERT_Output'))}, "
        f"start_offset={compact_value(values.get('start_offset'))}, "
        f"end_offset={compact_value(values.get('end_offset'))}"
    )


def describe_row_context(snapshot: dict[str, Any]) -> str:
    return (
        f"input_row={snapshot.get('input_excel_row')}, "
        f"output_row={snapshot.get('output_excel_row')}, "
        f"pmid={compact_value(snapshot.get('pmid'))}, "
        f"sent_index={compact_value(snapshot.get('sent_index'))}"
    )


def collect_group_changes(
    before_by_id: dict[int, dict[str, Any]],
    processed_rows: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    after_by_id = {id(row): row for row in processed_rows}
    changes: list[dict[str, Any]] = []

    for row_id, before in before_by_id.items():
        after_row = after_by_id.get(row_id)
        if after_row is None:
            changes.append({"action": "REMOVED", "before": before, "after_row": None})
            continue

        after = audit_snapshot(after_row)
        if tracked_values_changed(before, after):
            changes.append(
                {
                    "action": "MODIFIED",
                    "before": before,
                    "after_row": after_row,
                }
            )

    for row in processed_rows:
        if id(row) not in before_by_id:
            changes.append({"action": "ADDED", "before": None, "after_row": row})

    return changes


def hydrate_change_snapshots(changes: list[dict[str, Any]]) -> None:
    for change in changes:
        after_row = change.get("after_row")
        if after_row is not None:
            change["after"] = audit_snapshot(after_row)


def format_change(change: dict[str, Any]) -> str:
    action = change["action"]
    before = change.get("before")
    after = change.get("after")
    context = after or before

    if action == "ADDED":
        return f"[ADDED] {describe_row_context(after)}; new {describe_tracked_values(after)}"

    if action == "REMOVED":
        return f"[REMOVED] {describe_row_context(before)}; old {describe_tracked_values(before)}"

    return (
        f"[MODIFIED] {describe_row_context(context)}; "
        f"{describe_field_changes(before, after)}"
    )


def print_change_report(sheet_name: str, changes: list[dict[str, Any]]) -> None:
    if not changes:
        print(
            f"[{sheet_name}] Changed rows: 0 "
            "(ADDED=0, MODIFIED=0, REMOVED=0). No rows were added, modified, or removed."
        )
        return

    counts = Counter(change["action"] for change in changes)
    print(
        f"\n[{sheet_name}] Changed rows: {len(changes)} "
        f"(ADDED={counts['ADDED']}, MODIFIED={counts['MODIFIED']}, "
        f"REMOVED={counts['REMOVED']})"
    )

    if not PRINT_CHANGED_ROWS:
        return

    for index, change in enumerate(changes, start=1):
        if MAX_CHANGE_LINES is not None and index > MAX_CHANGE_LINES:
            print(f"... skipped {len(changes) - MAX_CHANGE_LINES} more changed rows")
            break
        print(format_change(change))


def term_text(row: dict[str, Any]) -> str:
    value = row.get("BERT_Output")
    return "" if is_blank(value) else str(value)


def sent_text(row: dict[str, Any]) -> str:
    value = row.get("sent_text")
    return "" if value is None else str(value)


def to_int(value: Any, default: int | None = None) -> int | None:
    if value is None or value == "":
        return default
    try:
        return int(value)
    except (TypeError, ValueError):
        return default


def valid_span(row: dict[str, Any]) -> tuple[int, int] | None:
    if is_blank(row.get("BERT_Output")):
        return None
    start = to_int(row.get("start_offset"))
    end = to_int(row.get("end_offset"))
    if start is None or end is None or start < 0 or end <= start:
        return None
    return start, end


def set_term(row: dict[str, Any], text: str | None, start: int = -1, end: int = -1) -> None:
    row["BERT_Output"] = text
    row["start_offset"] = start
    row["end_offset"] = end


def loose_normalize(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", value.casefold())


def normalize_spaces(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip().casefold()


def trim_match_span(text: str, start: int, end: int) -> tuple[str, int, int]:
    if start > 0 and text[start - 1] == "(" and end <= len(text) and text[end - 1] == ")":
        start -= 1
    while start < end and text[start].isspace():
        start += 1
    while end > start and text[end - 1].isspace():
        end -= 1
    return text[start:end], start, end


def spans_overlap(a_start: int, a_end: int, b_start: int, b_end: int) -> bool:
    return a_start < b_end and b_start < a_end


def row_sort_key(row: dict[str, Any]) -> tuple[int, int, int, int]:
    span = valid_span(row)
    if span is None:
        return (1, 10**12, 10**12, int(row.get("_source_order", 0)))
    start, end = span
    return (0, start, end, int(row.get("_source_order", 0)))


def make_new_row(
    base_row: dict[str, Any],
    bert_output: str,
    start: int,
    end: int,
) -> dict[str, Any]:
    new_row = copy.deepcopy(base_row)
    set_term(new_row, bert_output, start, end)
    return new_row


def is_available_in_existing_rows(
    candidate_text: str,
    candidate_start: int,
    candidate_end: int,
    rows: list[dict[str, Any]],
) -> bool:
    candidate_loose = loose_normalize(candidate_text)
    candidate_norm = normalize_spaces(candidate_text)
    if not candidate_loose:
        return True

    for row in rows:
        existing = term_text(row)
        if not existing:
            continue

        existing_loose = loose_normalize(existing)
        existing_norm = normalize_spaces(existing)
        if candidate_norm == existing_norm:
            return True
        if candidate_loose == existing_loose:
            return True
        if len(candidate_loose) >= 3 and candidate_loose in existing_loose:
            return True

        span = valid_span(row)
        if span == (candidate_start, candidate_end):
            return True

    return False


def remove_ca_125(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    ca_rows = [row for row in rows if CA_125_RE.search(term_text(row))]
    if not ca_rows:
        return rows

    other_term_rows = [
        row for row in rows if row not in ca_rows and not is_blank(row.get("BERT_Output"))
    ]
    if other_term_rows:
        return [row for row in rows if row not in ca_rows]

    keep = copy.deepcopy(ca_rows[0])
    set_term(keep, None, -1, -1)
    return [keep]


def merge_digit_comma_terms(rows: list[dict[str, Any]], text: str) -> list[dict[str, Any]]:
    active = sorted(rows, key=row_sort_key)
    consumed_ids: set[int] = set()

    for row in active:
        if id(row) in consumed_ids or not DIGIT_ONLY_RE.fullmatch(term_text(row)):
            continue

        span = valid_span(row)
        if span is None:
            continue
        start, end = span

        comma_match = re.match(r"\s*,\s*", text[end:])
        if not comma_match:
            continue

        next_start = end + comma_match.end()
        next_rows = []
        for candidate in active:
            if candidate is row or id(candidate) in consumed_ids:
                continue
            candidate_span = valid_span(candidate)
            if candidate_span and candidate_span[0] == next_start:
                next_rows.append(candidate)

        if not next_rows:
            continue

        next_row = sorted(next_rows, key=row_sort_key)[0]
        _, next_end = valid_span(next_row)  # type: ignore[misc]
        set_term(row, text[start:next_end], start, next_end)
        consumed_ids.add(id(next_row))

    return [row for row in rows if id(row) not in consumed_ids]


def merge_digit_o_tion_terms(rows: list[dict[str, Any]], text: str) -> list[dict[str, Any]]:
    active = sorted(rows, key=row_sort_key)
    consumed_ids: set[int] = set()

    for row in active:
        if id(row) in consumed_ids or not DIGIT_ONLY_RE.fullmatch(term_text(row)):
            continue

        span = valid_span(row)
        if span is None:
            continue
        start, end = span

        suffix_match = O_TION_AFTER_DIGIT_RE.match(text[end:])
        if not suffix_match:
            continue

        new_end = end + suffix_match.end()
        set_term(row, text[start:new_end], start, new_end)

        for candidate in active:
            if candidate is row:
                continue
            candidate_span = valid_span(candidate)
            if not candidate_span:
                continue
            c_start, c_end = candidate_span
            if c_start >= start and c_end <= new_end:
                consumed_ids.add(id(candidate))

    return [row for row in rows if id(row) not in consumed_ids]


def merge_alpha_beta_bridges(rows: list[dict[str, Any]], text: str) -> list[dict[str, Any]]:
    active = sorted(rows, key=row_sort_key)
    consumed_ids: set[int] = set()
    i = 0

    while i < len(active):
        first = active[i]
        if id(first) in consumed_ids or valid_span(first) is None:
            i += 1
            continue

        chain = [first]
        j = i
        while j + 1 < len(active):
            current = chain[-1]
            next_row = active[j + 1]
            if id(next_row) in consumed_ids:
                j += 1
                continue

            current_span = valid_span(current)
            next_span = valid_span(next_row)
            if current_span is None or next_span is None:
                break

            _, current_end = current_span
            next_start, _ = next_span
            if current_end > next_start:
                break

            gap = text[current_end:next_start]
            if not ALPHA_BETA_GAP_RE.fullmatch(gap):
                break

            chain.append(next_row)
            j += 1

        if len(chain) >= 2:
            start, _ = valid_span(chain[0])  # type: ignore[misc]
            _, end = valid_span(chain[-1])  # type: ignore[misc]
            set_term(chain[0], text[start:end], start, end)
            for merged_row in chain[1:]:
                consumed_ids.add(id(merged_row))
            i = j + 1
        else:
            i += 1

    return [row for row in rows if id(row) not in consumed_ids]


def add_regex_pattern_rows(
    rows: list[dict[str, Any]],
    text: str,
    regexes: list[re.Pattern[str]],
) -> list[dict[str, Any]]:
    if not text:
        return rows

    candidates: list[tuple[int, int, str]] = []
    seen_candidates: set[tuple[int, int, str]] = set()

    for regex in regexes:
        for match in regex.finditer(text):
            term, start, end = trim_match_span(text, match.start(), match.end())
            if not term:
                continue
            key = (start, end, loose_normalize(term))
            if key in seen_candidates:
                continue
            seen_candidates.add(key)
            candidates.append((start, end, term))

    candidates.sort(key=lambda item: (item[0], -(item[1] - item[0]), item[2].casefold()))

    base_row = rows[0] if rows else {"sent_text": text}
    new_rows: list[dict[str, Any]] = []
    occupied_new_spans: list[tuple[int, int]] = []

    for start, end, term in candidates:
        if is_available_in_existing_rows(term, start, end, rows + new_rows):
            continue
        if any(spans_overlap(start, end, s, e) for s, e in occupied_new_spans):
            continue

        new_row = make_new_row(base_row, term, start, end)
        new_rows.append(new_row)
        occupied_new_spans.append((start, end))

    return rows + new_rows


def add_lewis_pattern_rows(rows: list[dict[str, Any]], text: str) -> list[dict[str, Any]]:
    return add_regex_pattern_rows(rows, text, LEWIS_REGEXES)


def add_additional_term_rows(rows: list[dict[str, Any]], text: str) -> list[dict[str, Any]]:
    return add_regex_pattern_rows(rows, text, ADDITIONAL_TERM_REGEXES)


def merge_single_space_adjacent_terms(rows: list[dict[str, Any]], text: str) -> list[dict[str, Any]]:
    active = sorted(rows, key=row_sort_key)
    consumed_ids: set[int] = set()

    for row in active:
        if id(row) in consumed_ids:
            continue

        span = valid_span(row)
        if span is None:
            continue

        start, current_end = span
        changed = False

        while True:
            next_rows = []
            for candidate in active:
                if candidate is row or id(candidate) in consumed_ids:
                    continue
                candidate_span = valid_span(candidate)
                if not candidate_span:
                    continue
                next_start, next_end = candidate_span
                if text[current_end:next_start] == " ":
                    next_rows.append((next_start, next_end, candidate))

            if not next_rows:
                break

            _, current_end, next_row = sorted(next_rows, key=lambda item: (item[0], item[1]))[0]
            consumed_ids.add(id(next_row))
            changed = True

        if changed:
            set_term(row, text[start:current_end], start, current_end)

    return [row for row in rows if id(row) not in consumed_ids]


def drop_blank_rows_when_terms_exist(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    has_terms = any(not is_blank(row.get("BERT_Output")) for row in rows)
    if not has_terms:
        return rows[:1]
    return [row for row in rows if not is_blank(row.get("BERT_Output"))]


def process_group(rows: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    text = sent_text(rows[0]) if rows else ""
    before_by_id = {id(row): audit_snapshot(row) for row in rows}

    rows = remove_ca_125(rows)
    rows = merge_digit_comma_terms(rows, text)
    rows = merge_digit_o_tion_terms(rows, text)
    rows = merge_alpha_beta_bridges(rows, text)
    rows = add_lewis_pattern_rows(rows, text)
    rows = add_additional_term_rows(rows, text)
    rows = drop_blank_rows_when_terms_exist(rows)
    rows = merge_single_space_adjacent_terms(rows, text)

    processed_rows = sorted(rows, key=row_sort_key)
    return processed_rows, collect_group_changes(before_by_id, processed_rows)


def process_rows(rows: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    grouped: OrderedDict[tuple[Any, Any, Any], list[dict[str, Any]]] = OrderedDict()

    for index, row in enumerate(rows):
        row["_source_order"] = index
        key = (row.get("pmid"), row.get("sent_index"), row.get("sent_text"))
        grouped.setdefault(key, []).append(row)

    processed: list[dict[str, Any]] = []
    changes: list[dict[str, Any]] = []
    for group_rows in grouped.values():
        group_processed, group_changes = process_group(group_rows)
        processed.extend(group_processed)
        changes.extend(group_changes)

    for output_excel_row, row in enumerate(processed, start=2):
        row["_output_excel_row"] = output_excel_row

    hydrate_change_snapshots(changes)

    return processed, changes


def worksheet_has_required_columns(headers: list[Any]) -> bool:
    header_set = {str(header) for header in headers if header is not None}
    return all(column in header_set for column in REQUIRED_COLUMNS)


def read_sheet_as_dicts(ws: openpyxl.worksheet.worksheet.Worksheet) -> tuple[list[str], list[dict[str, Any]]]:
    raw_headers = [cell.value for cell in ws[1]]
    headers = [str(header) if header is not None else "" for header in raw_headers]
    rows: list[dict[str, Any]] = []

    for excel_row, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        row = {header: value for header, value in zip(headers, values)}
        row["_excel_row"] = excel_row
        rows.append(row)

    return headers, rows


def write_sheet_from_dicts(
    ws: openpyxl.worksheet.worksheet.Worksheet,
    headers: list[str],
    rows: list[dict[str, Any]],
) -> None:
    for col_index, header in enumerate(headers, start=1):
        ws.cell(row=1, column=col_index, value=header)

    for row_index, row_data in enumerate(rows, start=2):
        for col_index, header in enumerate(headers, start=1):
            ws.cell(row=row_index, column=col_index, value=row_data.get(header))


def copy_sheet_values(
    source_ws: openpyxl.worksheet.worksheet.Worksheet,
    target_ws: openpyxl.worksheet.worksheet.Worksheet,
) -> None:
    for row in source_ws.iter_rows(values_only=True):
        target_ws.append(row)


def process_workbook(input_path: Path, output_path: Path) -> None:
    source_wb = openpyxl.load_workbook(input_path, data_only=False)
    output_wb = openpyxl.Workbook()
    output_wb.remove(output_wb.active)

    for source_ws in source_wb.worksheets:
        target_ws = output_wb.create_sheet(source_ws.title)
        headers, rows = read_sheet_as_dicts(source_ws)

        if worksheet_has_required_columns(headers):
            output_headers = [
                header for header in headers if header not in REMOVED_OUTPUT_COLUMNS
            ]
            processed_rows, changes = process_rows(rows)
            write_sheet_from_dicts(target_ws, output_headers, processed_rows)
            print_change_report(source_ws.title, changes)
        else:
            copy_sheet_values(source_ws, target_ws)
            print(f"[{source_ws.title}] Skipped change report; required columns not found.")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_wb.save(output_path)


def main() -> None:
    process_workbook(INPUT_FILE, OUTPUT_FILE)
    print(f"Saved updated workbook to: {OUTPUT_FILE}")


if __name__ == "__main__":
    main()
