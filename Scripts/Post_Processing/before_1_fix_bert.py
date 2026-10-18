#!/usr/bin/env python3
"""
Clean BERT_Output rows in an Excel workbook.

Rules implemented:
1. Remove trailing punctuation/hyphens from BERT_Output values.
2. If a value has an opening parenthesis but no closing parenthesis and the
   source sent_text has the missing ")" immediately after the entity, add it.
3. Update start_offset and end_offset for every changed BERT_Output.
4. Add/update a Modify column and mark "Yes" when BERT_Output changed.

Offsets are treated as 0-based, with end_offset exclusive.
"""

from __future__ import annotations
from glycaner_paths import STAGE3_INPUT, resource, work  # noqa: E402

from copy import copy
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell


REQUIRED_COLUMNS = ["sent_text", "BERT_Output", "start_offset", "end_offset"]
MODIFY_COLUMN = "Modify"

# Put your input and output files here.
INPUT_FILE = Path(work("bert_output_full_updated.xlsx"))
OUTPUT_FILE = Path(work("bert_output_full_updated_modified.xlsx"))
SHEET_NAME = None  # Use None for the active sheet, or set this to "Sheet1".

# Do not include ")" here because a valid entity may end with a balanced
# closing parenthesis, such as sialyl-Lewis(x).
TRAILING_CHARS = " \t\r\n,;:.!?-–—/\\|[]{}\"'"


@dataclass
class Segment:
    text: str
    start: int
    end: int


def normalize_header(value: Any) -> str:
    return str(value).strip().lower() if value is not None else ""


def get_column_indexes(ws) -> dict[str, int]:
    headers = {
        normalize_header(ws.cell(row=1, column=col).value): col
        for col in range(1, ws.max_column + 1)
    }

    missing = [name for name in REQUIRED_COLUMNS if normalize_header(name) not in headers]
    if missing:
        raise ValueError(f"Missing required column(s): {', '.join(missing)}")

    if normalize_header(MODIFY_COLUMN) not in headers:
        modify_col = ws.max_column + 1
        ws.cell(row=1, column=modify_col).value = MODIFY_COLUMN

        # Copy header styling from the previous column, if present.
        if modify_col > 1:
            copy_cell_format(ws.cell(row=1, column=modify_col - 1), ws.cell(row=1, column=modify_col))
    else:
        modify_col = headers[normalize_header(MODIFY_COLUMN)]

    headers[normalize_header(MODIFY_COLUMN)] = modify_col
    return headers


def copy_cell_format(source: Cell, target: Cell) -> None:
    if source.has_style:
        target.font = copy(source.font)
        target.fill = copy(source.fill)
        target.border = copy(source.border)
        target.alignment = copy(source.alignment)
        target.protection = copy(source.protection)
        target.number_format = source.number_format
    if source.hyperlink:
        target._hyperlink = copy(source.hyperlink)
    if source.comment:
        target.comment = copy(source.comment)


def read_row_template(ws, row_number: int, max_col: int) -> list[dict[str, Any]]:
    template = []
    for col in range(1, max_col + 1):
        cell = ws.cell(row=row_number, column=col)
        template.append(
            {
                "value": cell.value,
                "font": copy(cell.font),
                "fill": copy(cell.fill),
                "border": copy(cell.border),
                "alignment": copy(cell.alignment),
                "protection": copy(cell.protection),
                "number_format": cell.number_format,
                "hyperlink": copy(cell.hyperlink) if cell.hyperlink else None,
                "comment": copy(cell.comment) if cell.comment else None,
                "has_style": cell.has_style,
            }
        )
    return template


def write_row_from_template(ws, row_number: int, template: list[dict[str, Any]]) -> None:
    for col, info in enumerate(template, start=1):
        cell = ws.cell(row=row_number, column=col)
        cell.value = info["value"]
        if info["has_style"]:
            cell.font = copy(info["font"])
            cell.fill = copy(info["fill"])
            cell.border = copy(info["border"])
            cell.alignment = copy(info["alignment"])
            cell.protection = copy(info["protection"])
            cell.number_format = info["number_format"]
        if info["hyperlink"]:
            cell._hyperlink = copy(info["hyperlink"])
        if info["comment"]:
            cell.comment = copy(info["comment"])


def as_int(value: Any) -> int | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    try:
        return int(str(value).strip())
    except (TypeError, ValueError):
        return None


def find_span(sent_text: str, bert_output: str, start_offset: Any, end_offset: Any) -> tuple[int | None, int | None]:
    """Find where bert_output lives in sent_text, preferring the existing offsets."""
    start = as_int(start_offset)
    end = as_int(end_offset)

    if start is not None and end is not None and 0 <= start <= end <= len(sent_text):
        if sent_text[start:end] == bert_output:
            return start, end
        if end - start == len(bert_output):
            return start, end

    found = sent_text.find(bert_output)
    if found != -1:
        return found, found + len(bert_output)

    lower_found = sent_text.lower().find(bert_output.lower())
    if lower_found != -1:
        return lower_found, lower_found + len(bert_output)

    return start, end


def trim_trailing(text: str, absolute_start: int | None) -> tuple[str, int | None, int]:
    """Trim trailing punctuation/hyphens and return text, start, number of chars removed."""
    original_len = len(text)
    cleaned = text.rstrip(TRAILING_CHARS)

    # Remove stray closing parentheses only when they are not balanced.
    while cleaned.endswith(")") and cleaned.count(")") > cleaned.count("("):
        cleaned = cleaned[:-1].rstrip(TRAILING_CHARS)

    removed = original_len - len(cleaned)

    if absolute_start is None:
        return cleaned, None, removed
    return cleaned, absolute_start, removed


def add_missing_closing_parenthesis(
    text: str,
    sent_text: str,
    absolute_end: int | None,
) -> tuple[str, int | None, bool]:
    if not text:
        return text, absolute_end, False
    if text.count("(") <= text.count(")"):
        return text, absolute_end, False
    if not text[-1].isalpha():
        return text, absolute_end, False
    if absolute_end is None or absolute_end < 0 or absolute_end >= len(sent_text):
        return text, absolute_end, False
    if sent_text[absolute_end] != ")":
        return text, absolute_end, False
    return text + ")", absolute_end + 1, True


def clean_bert_output(
    bert_output: str,
    sent_text: str,
    start_offset: Any,
    end_offset: Any,
) -> list[Segment]:
    base_start, base_end = find_span(sent_text, bert_output, start_offset, end_offset)
    segments: list[Segment] = []

    left_trimmed = bert_output.lstrip()
    leading_removed = len(bert_output) - len(left_trimmed)

    if base_start is None:
        absolute_start = None
    else:
        absolute_start = base_start + leading_removed

    cleaned, absolute_start, trailing_removed = trim_trailing(left_trimmed, absolute_start)
    if not cleaned:
        return segments

    if absolute_start is None:
        absolute_end = None
    else:
        absolute_end = absolute_start + len(left_trimmed) - trailing_removed

    cleaned, absolute_end, _ = add_missing_closing_parenthesis(cleaned, sent_text, absolute_end)

    if absolute_start is None or absolute_end is None:
        absolute_start = -1
        absolute_end = -1

    segments.append(Segment(cleaned, absolute_start, absolute_end))

    return segments


def process_workbook(input_path: Path, output_path: Path, sheet_name: str | None = None) -> dict[str, int]:
    wb = load_workbook(input_path)
    ws = wb[sheet_name] if sheet_name else wb.active

    cols = get_column_indexes(ws)
    max_col = ws.max_column

    bert_col = cols[normalize_header("BERT_Output")]
    sent_col = cols[normalize_header("sent_text")]
    start_col = cols[normalize_header("start_offset")]
    end_col = cols[normalize_header("end_offset")]
    modify_col = cols[normalize_header(MODIFY_COLUMN)]

    original_rows = ws.max_row - 1
    output_templates: list[list[dict[str, Any]]] = []
    modified_rows = 0

    for row_number in range(2, ws.max_row + 1):
        template = read_row_template(ws, row_number, max_col)

        bert_output = template[bert_col - 1]["value"]
        sent_text = template[sent_col - 1]["value"]
        start_offset = template[start_col - 1]["value"]
        end_offset = template[end_col - 1]["value"]

        if not isinstance(bert_output, str) or not bert_output.strip() or not isinstance(sent_text, str):
            template[modify_col - 1]["value"] = None
            output_templates.append(template)
            continue

        segments = clean_bert_output(bert_output, sent_text, start_offset, end_offset)

        if not segments:
            template[modify_col - 1]["value"] = None
            output_templates.append(template)
            continue

        changed = len(segments) != 1 or segments[0].text != bert_output

        if not changed:
            template[modify_col - 1]["value"] = None
            output_templates.append(template)
            continue

        modified_rows += 1

        for segment in segments:
            new_template = [
                {
                    key: copy(value)
                    for key, value in cell_info.items()
                }
                for cell_info in template
            ]
            new_template[bert_col - 1]["value"] = segment.text
            new_template[start_col - 1]["value"] = segment.start
            new_template[end_col - 1]["value"] = segment.end
            new_template[modify_col - 1]["value"] = "Yes"
            output_templates.append(new_template)

    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)

    for row_index, template in enumerate(output_templates, start=2):
        write_row_from_template(ws, row_index, template)

    # Keep the filter range, if any, aligned with the new sheet size.
    if ws.auto_filter and ws.auto_filter.ref:
        ws.auto_filter.ref = ws.dimensions

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)

    return {
        "original_rows": original_rows,
        "output_rows": len(output_templates),
        "modified_original_rows": modified_rows,
    }


def main() -> None:
    input_path = INPUT_FILE.expanduser().resolve()
    output_path = OUTPUT_FILE.expanduser().resolve()

    summary = process_workbook(input_path, output_path, SHEET_NAME)

    print(f"Saved: {output_path}")
    print(f"Original data rows: {summary['original_rows']}")
    print(f"Output data rows: {summary['output_rows']}")
    print(f"Original rows modified: {summary['modified_original_rows']}")


if __name__ == "__main__":
    main()
