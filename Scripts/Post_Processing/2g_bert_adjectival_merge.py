#!/usr/bin/env python3
"""
Merge pattern-bearing BERT rows connected as comma/conjunction lists.

For each pmid and sent_index, rows are ordered by their positions in sent_text.
A sequence is merged only when one or more incomplete pattern elements are
completed by the final element, their gaps contain only list punctuation, and
the gap immediately before the final element contains "and" or "or". The final
element does not have to match one of the target patterns.

The retained row receives exact sentence substrings in BERT_Output and
BERT_Base_NPs, updated BERT_Output offsets, and 2g="Yes". The BaseNP may begin
before BERT_Output (for example, "the pure 6-sulfo"). Other rows in the merged
sequence are removed.
"""

from __future__ import annotations
from glycaner_paths import STAGE3_INPUT, resource, work  # noqa: E402

import re
from copy import copy
from pathlib import Path

from openpyxl import load_workbook


INPUT_FILE = Path(
    work("2f_addition_and_or_updated.xlsx")
)
OUTPUT_FILE = INPUT_FILE.with_name("2g_pattern_list_merged.xlsx")
SHEET_NAME = None  # use None to process every worksheet with required columns

PMID_COL = "pmid"
SENT_INDEX_COL = "sent_index"
SENT_TEXT_COL = "sent_text"
BERT_OUTPUT_COL = "BERT_Output"
START_COL = "start_offset"
END_COL = "end_offset"
BERT_BASE_COL = "BERT_Base_NPs"
DISCARD_COL = "Discard"
FLAG_COL = "2g"

REQUIRED_COLUMNS = {
    PMID_COL,
    SENT_INDEX_COL,
    SENT_TEXT_COL,
    BERT_OUTPUT_COL,
    START_COL,
    END_COL,
    BERT_BASE_COL,
}

PATTERN_STRINGS = [
    r"\b\w*osyl\b",
    r"\b\w*aminyl\b",
    r"\b\w*lyl\b",
    r"\bmethyl\b",
    r"\b\w*sulfo\b",
    r"\btype[-\s]?[1-6](?:/[1-6])*\b",
    r"\btypes?\s+(?:A|B|O|H)[-\s]+(?:[1-4]|I{1,3}|IV)\b",
    r"\bbi\b",
    r"\bdi\b",
    r"\bdim\b",
    r"\btri\b",
    r"\btria\b",
    r"\btetra\b",
    r"\bmono\b",
    r"\bpara\b",
    r"\bpenta\b",
    r"\bhexa\b",
    r"\bhepta\b",
    r"\bpoly\b",
    r"\boligo\b",
    r"\blyso\b",
    r"\bocta\b",
    r"\bnona\b",
    r"\bneo\b",
    r"\blacto\b",
    r"\bgluco\b",
    r"\bglobo\b",
    r"\bgalacto\b",
    r"\bganglio\b",
    r"\bcore[-\s]?[A-Za-z]?\d+\b",
    r"\bcore[-\s]?(?:[IVX]+)\b",
    r"\b\w+ed\b",
]
TARGET_REGEXES = [re.compile(pattern, re.IGNORECASE) for pattern in PATTERN_STRINGS]

# Generic head nouns that should not be carried into the merged list. They are
# removed only when they occur at the very end of the final list component.
MERGED_TRAILING_HEAD_RE = re.compile(r"\s+determinants?$", re.IGNORECASE)

# Comma-list gaps may contain punctuation only. A merge is emitted only when
# at least one gap also contains an explicit and/or connector.
PUNCTUATION_GAP_RE = re.compile(r"^[\s,;:/\\|()\[\]{}'\"\-\u2013\u2014]*$")
CONJUNCTION_GAP_RE = re.compile(
    r"^[\s,;:/\\|()\[\]{}'\"\-\u2013\u2014]*"
    r"(?:and|or)"
    r"(?:\s+(?:a|an|the))?"
    r"[\s,;:/\\|()\[\]{}'\"\-\u2013\u2014]*$",
    re.IGNORECASE,
)


def normalize_spaces(value) -> str:
    return re.sub(r"\s+", " ", str(value).strip())


def cell_text(value) -> str:
    if value is None:
        return ""
    return normalize_spaces(value)


def is_yes(value) -> bool:
    return str(value).strip().casefold() == "yes"


def to_int(value):
    try:
        return int(value)
    except (TypeError, ValueError):
        return None


def get_header_map(ws):
    return {
        str(cell.value).strip(): cell.column
        for cell in ws[1]
        if cell.value is not None
    }


def copy_cell_style(source, target) -> None:
    if source.has_style:
        target._style = copy(source._style)
    target.font = copy(source.font)
    target.fill = copy(source.fill)
    target.border = copy(source.border)
    target.alignment = copy(source.alignment)
    target.number_format = source.number_format
    target.protection = copy(source.protection)


def ensure_column(ws, headers, column_name: str):
    if column_name in headers:
        return headers

    column = ws.max_column + 1
    ws.cell(row=1, column=column).value = column_name
    if column > 1:
        copy_cell_style(
            ws.cell(row=1, column=column - 1),
            ws.cell(row=1, column=column),
        )
    return get_header_map(ws)


def valid_output_span(ws, row_idx: int, headers):
    start = to_int(ws.cell(row=row_idx, column=headers[START_COL]).value)
    end = to_int(ws.cell(row=row_idx, column=headers[END_COL]).value)
    if start is None or end is None or start < 0 or end <= start:
        return None
    return start, end


def find_spans(text: str, value: str):
    if not text or not value:
        return []

    text_folded = text.casefold()
    value_folded = value.casefold()
    spans = []
    start = 0

    while True:
        position = text_folded.find(value_folded, start)
        if position < 0:
            break
        spans.append((position, position + len(value)))
        start = position + 1

    return spans


def overlap_length(first, second) -> int:
    return max(0, min(first[1], second[1]) - max(first[0], second[0]))


def choose_text_span(sent_text: str, value: str, output_span=None):
    candidates = find_spans(sent_text, cell_text(value))
    if not candidates:
        return output_span
    if output_span is None:
        return candidates[0]

    best = max(candidates, key=lambda span: overlap_length(span, output_span))
    if overlap_length(best, output_span) > 0:
        return best

    return min(
        candidates,
        key=lambda span: (
            abs(span[0] - output_span[0]) + abs(span[1] - output_span[1]),
            span[0],
        ),
    )


def value_is_incomplete_pattern(value: str) -> bool:
    """Return True when a pattern has no meaningful completion after it."""
    for regex in TARGET_REGEXES:
        for match in regex.finditer(value):
            suffix_words = re.findall(r"[A-Za-z]+", value[match.end() :])

            # Ignore one-letter biochemical modifiers such as O-, N-, or A-.
            meaningful_words = [word for word in suffix_words if len(word) > 1]
            if not meaningful_words:
                return True

    return False


def component_is_incomplete(component) -> bool:
    return value_is_incomplete_pattern(
        component["bert_output"]
    ) and value_is_incomplete_pattern(component["bert_base"])


def punctuation_only_gap(gap: str) -> bool:
    return bool(PUNCTUATION_GAP_RE.fullmatch(gap))


def conjunction_gap(gap: str) -> bool:
    return bool(CONJUNCTION_GAP_RE.fullmatch(gap))


def row_group_key(ws, row_idx: int, headers):
    return (
        ws.cell(row=row_idx, column=headers[PMID_COL]).value,
        ws.cell(row=row_idx, column=headers[SENT_INDEX_COL]).value,
        ws.cell(row=row_idx, column=headers[SENT_TEXT_COL]).value,
    )


def build_components(ws, row_indices, headers):
    components = []

    for row_idx in row_indices:
        if DISCARD_COL in headers and is_yes(
            ws.cell(row=row_idx, column=headers[DISCARD_COL]).value
        ):
            continue

        bert_output = cell_text(
            ws.cell(row=row_idx, column=headers[BERT_OUTPUT_COL]).value
        )
        bert_base = cell_text(
            ws.cell(row=row_idx, column=headers[BERT_BASE_COL]).value
        )
        sent_value = ws.cell(row=row_idx, column=headers[SENT_TEXT_COL]).value
        sent_text = "" if sent_value is None else str(sent_value)
        output_span = valid_output_span(ws, row_idx, headers)

        if (
            not sent_text
            or not bert_output
            or not bert_base
            or output_span is None
        ):
            continue

        base_span = choose_text_span(sent_text, bert_base, output_span)
        if base_span is None:
            base_span = output_span

        components.append(
            {
                "row_idx": row_idx,
                "sent_text": sent_text,
                "bert_output": bert_output,
                "bert_base": bert_base,
                "output_span": output_span,
                "base_span": base_span,
            }
        )

    return sorted(
        components,
        key=lambda item: (
            item["base_span"][0],
            item["base_span"][1],
            item["output_span"][0],
            item["row_idx"],
        ),
    )


def finish_sequence(sequences, sequence, has_conjunction) -> None:
    if len(sequence) >= 2 and has_conjunction:
        sequences.append(sequence)


def find_merge_sequences(components):
    sequences = []
    if len(components) < 2:
        return sequences

    current = [components[0]]
    current_has_conjunction = False

    for previous, current_item in zip(components, components[1:]):
        previous_end = previous["base_span"][1]
        current_start = current_item["base_span"][0]

        if current_start < previous_end:
            finish_sequence(sequences, current, current_has_conjunction)
            current = [current_item]
            current_has_conjunction = False
            continue

        gap = previous["sent_text"][previous_end:current_start]
        if punctuation_only_gap(gap) or conjunction_gap(gap):
            current.append(current_item)
            current_has_conjunction = (
                current_has_conjunction or conjunction_gap(gap)
            )
            continue

        finish_sequence(sequences, current, current_has_conjunction)
        current = [current_item]
        current_has_conjunction = False

    finish_sequence(sequences, current, current_has_conjunction)
    return sequences


def final_gap_has_conjunction(sequence) -> bool:
    if len(sequence) < 2:
        return False

    previous = sequence[-2]
    final = sequence[-1]
    previous_end = previous["base_span"][1]
    final_start = final["base_span"][0]
    if final_start < previous_end:
        return False

    gap = previous["sent_text"][previous_end:final_start]
    return conjunction_gap(gap)


def sequence_has_parenthetical_or_percentage(sequence) -> bool:
    sent_text = sequence[0]["sent_text"]
    start = min(item["output_span"][0] for item in sequence)
    end = max(
        max(item["output_span"][1], item["base_span"][1])
        for item in sequence
    )
    merged_region = sent_text[start:end]
    return bool(re.search(r"[()%]", merged_region))


def incomplete_element_sequences(sequence):
    """Extract incomplete-prefix runs ending in one complete final element."""
    results = []
    incomplete_run = []

    for component in sequence:
        if component_is_incomplete(component):
            incomplete_run.append(component)
            continue

        if incomplete_run:
            candidate = incomplete_run + [component]
            if (
                len(candidate) >= 2
                and final_gap_has_conjunction(candidate)
                and not sequence_has_parenthetical_or_percentage(candidate)
            ):
                results.append(candidate)

        incomplete_run = []

    return results


def trim_merged_span(sent_text: str, start: int, end: int):
    while start < end and sent_text[start].isspace():
        start += 1
    while end > start and sent_text[end - 1].isspace():
        end -= 1

    # Do not retain sentence/list punctuation after the final term.
    while end > start and sent_text[end - 1] in ",;:/|\u2013\u2014":
        end -= 1
        while end > start and sent_text[end - 1].isspace():
            end -= 1

    return (start, end) if start < end else None


def trim_trailing_head(sent_text: str, start: int, end: int):
    match = MERGED_TRAILING_HEAD_RE.search(sent_text[start:end])
    if match is not None:
        end = start + match.start()
    return (start, end) if start < end else None


def sequence_values(sequence):
    sent_text = sequence[0]["sent_text"]
    output_start = min(item["output_span"][0] for item in sequence)
    base_start = min(item["base_span"][0] for item in sequence)
    end = max(
        max(item["output_span"][1], item["base_span"][1])
        for item in sequence
    )
    output_span = trim_merged_span(sent_text, output_start, end)
    base_span = trim_merged_span(sent_text, base_start, end)
    if output_span is None or base_span is None:
        return None

    output_span = trim_trailing_head(sent_text, *output_span)
    base_span = trim_trailing_head(sent_text, *base_span)
    if output_span is None or base_span is None:
        return None

    output_start, output_end = output_span
    base_start, base_end = base_span
    return {
        "bert_output": sent_text[output_start:output_end],
        "bert_base": sent_text[base_start:base_end],
        "start": output_start,
        "end": output_end,
    }


def sequence_key(group_key, bert_output: str, bert_base: str):
    pmid, sent_index, sent_text = group_key
    return (
        pmid,
        sent_index,
        sent_text,
        normalize_spaces(bert_output).casefold(),
        normalize_spaces(bert_base).casefold(),
    )


def existing_keys(ws, headers):
    keys = set()
    for row_idx in range(2, ws.max_row + 1):
        bert_output = cell_text(
            ws.cell(row=row_idx, column=headers[BERT_OUTPUT_COL]).value
        )
        bert_base = cell_text(
            ws.cell(row=row_idx, column=headers[BERT_BASE_COL]).value
        )
        if not bert_output or not bert_base:
            continue

        group_key = row_group_key(ws, row_idx, headers)
        keys.add(
            (
                group_key[0],
                group_key[1],
                group_key[2],
                normalize_spaces(bert_output).casefold(),
                normalize_spaces(bert_base).casefold(),
            )
        )
    return keys


def process_worksheet(ws) -> dict[str, int]:
    headers = get_header_map(ws)
    if not REQUIRED_COLUMNS.issubset(headers):
        return {
            "input_rows": ws.max_row - 1,
            "rows_replaced": 0,
            "rows_removed": 0,
            "output_rows": ws.max_row - 1,
        }

    headers = ensure_column(ws, headers, FLAG_COL)

    groups = {}
    for row_idx in range(2, ws.max_row + 1):
        groups.setdefault(row_group_key(ws, row_idx, headers), []).append(row_idx)

    known_keys = existing_keys(ws, headers)
    row_updates = {}
    rows_to_delete = set()
    claimed_rows = set()

    for group_key, row_indices in groups.items():
        components = build_components(ws, row_indices, headers)
        raw_sequences = find_merge_sequences(components)
        sequences = [
            candidate
            for raw_sequence in raw_sequences
            for candidate in incomplete_element_sequences(raw_sequence)
        ]

        for sequence in sequences:
            sequence_rows = {item["row_idx"] for item in sequence}
            if sequence_rows & claimed_rows:
                continue

            values = sequence_values(sequence)
            if values is None:
                continue

            key = sequence_key(
                group_key,
                values["bert_output"],
                values["bert_base"],
            )
            if key in known_keys:
                continue
            known_keys.add(key)

            keep_row = sequence[0]["row_idx"]
            row_updates[keep_row] = values
            rows_to_delete.update(sequence_rows - {keep_row})
            claimed_rows.update(sequence_rows)

    input_rows = ws.max_row - 1
    for row_idx, update in row_updates.items():
        ws.cell(row=row_idx, column=headers[BERT_OUTPUT_COL]).value = update[
            "bert_output"
        ]
        ws.cell(row=row_idx, column=headers[BERT_BASE_COL]).value = update[
            "bert_base"
        ]
        ws.cell(row=row_idx, column=headers[START_COL]).value = update["start"]
        ws.cell(row=row_idx, column=headers[END_COL]).value = update["end"]
        ws.cell(row=row_idx, column=headers[FLAG_COL]).value = "Yes"

    for row_idx in sorted(rows_to_delete, reverse=True):
        ws.delete_rows(row_idx, 1)

    return {
        "input_rows": input_rows,
        "rows_replaced": len(row_updates),
        "rows_removed": len(rows_to_delete),
        "output_rows": ws.max_row - 1,
    }


def process_workbook(input_path: Path, output_path: Path) -> dict[str, int]:
    if not input_path.exists():
        raise FileNotFoundError(f"Input file not found: {input_path}")

    workbook = load_workbook(input_path)
    totals = {
        "input_rows": 0,
        "rows_replaced": 0,
        "rows_removed": 0,
        "output_rows": 0,
    }

    worksheets = (
        [workbook[SHEET_NAME]]
        if SHEET_NAME is not None
        else workbook.worksheets
    )
    for worksheet in worksheets:
        summary = process_worksheet(worksheet)
        for key, value in summary.items():
            totals[key] += value

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return totals


def main() -> None:
    summary = process_workbook(INPUT_FILE, OUTPUT_FILE)

    print(f"Input rows: {summary['input_rows']}")
    print(f"Rows replaced: {summary['rows_replaced']}")
    print(f"Rows removed: {summary['rows_removed']}")
    print(f"Output rows: {summary['output_rows']}")
    print(f"Saved output file to: {OUTPUT_FILE}")


if __name__ == "__main__":
    main()
