#!/usr/bin/env python3
"""
Combine eligible type/core numeric BERT rows connected by "and" or "or".

The workbook contains BERT_Output, BERT_Base_NPs, and Addition. This script
updates only BERT_Output, BERT_Base_NPs, start_offset, end_offset, and Addition.
It does not require or create Merge or Merged_Base_NPs columns.
"""

from __future__ import annotations
from glycaner_paths import STAGE3_INPUT, resource, work  # noqa: E402

import re
from copy import copy
from pathlib import Path

from openpyxl import load_workbook


INPUT_FILE = Path(
    work("2e_nested_base_np_updated.xlsx")
)
OUTPUT_FILE = INPUT_FILE.with_name("2f_addition_and_or_updated.xlsx")
SHEET_NAME = None

PMID_COL = "pmid"
SENT_INDEX_COL = "sent_index"
SENT_TEXT_COL = "sent_text"
BERT_OUTPUT_COL = "BERT_Output"
START_COL = "start_offset"
END_COL = "end_offset"
BERT_BASE_COL = "BERT_Base_NPs"
DISCARD_COL = "Discard"
ADDITION_COL = "Addition"

REQUIRED_COLUMNS = [
    PMID_COL,
    SENT_INDEX_COL,
    SENT_TEXT_COL,
    BERT_OUTPUT_COL,
    START_COL,
    END_COL,
    BERT_BASE_COL,
]

CONNECTOR_GAP_RE = re.compile(
    r"^[\s,;:/0-9-]*(?:and|or)(?:\s+(?:a|an|the))?[\s,;:/0-9-]*$",
    re.IGNORECASE,
)
LIST_EXTENSION_GAP_RE = re.compile(r"^[\s,;:/0-9-]*$")
CODED_NUMBER_PATTERN = r"(?:[A-Za-z]\s*-?)?-?\d+"
GAP_DIGIT_BEFORE_AND_OR_RE = re.compile(
    rf"{CODED_NUMBER_PATTERN}[\s,;:/0-9A-Za-z()\[\]-]*\b(?:and|or)\b",
    re.IGNORECASE,
)
TYPE_CORE_NUMBER_RE = re.compile(
    rf"\b(?:types?|cores?)\s*-?{CODED_NUMBER_PATTERN}\b",
    re.IGNORECASE,
)
LEADING_NUMBER_RE = re.compile(
    rf"^\s*{CODED_NUMBER_PATTERN}\b",
    re.IGNORECASE,
)


def normalize_spaces(value) -> str:
    return re.sub(r"\s+", " ", str(value).strip())


def cell_text(value) -> str:
    if value is None:
        return ""
    return normalize_spaces(value)


def is_yes(value) -> bool:
    return str(value).strip().casefold() == "yes"


def to_int(value):
    try:
        return int(value)
    except (TypeError, ValueError):
        return None


def get_header_map(ws):
    return {
        str(cell.value).strip(): cell.column
        for cell in ws[1]
        if cell.value is not None
    }


def ensure_column(ws, headers, column_name: str):
    if column_name in headers:
        return headers

    col_idx = ws.max_column + 1
    ws.cell(row=1, column=col_idx).value = column_name

    if col_idx > 1:
        source = ws.cell(row=1, column=col_idx - 1)
        target = ws.cell(row=1, column=col_idx)
        copy_cell_style(source, target)

    return get_header_map(ws)


def copy_cell_style(source, target) -> None:
    if source.has_style:
        target._style = copy(source._style)
    target.font = copy(source.font)
    target.fill = copy(source.fill)
    target.border = copy(source.border)
    target.alignment = copy(source.alignment)
    target.number_format = source.number_format
    target.protection = copy(source.protection)


def copy_row(ws, source_row: int, target_row: int) -> None:
    for col_idx in range(1, ws.max_column + 1):
        source = ws.cell(row=source_row, column=col_idx)
        target = ws.cell(row=target_row, column=col_idx)
        target.value = source.value
        copy_cell_style(source, target)

    source_dim = ws.row_dimensions[source_row]
    target_dim = ws.row_dimensions[target_row]
    target_dim.height = source_dim.height
    target_dim.hidden = source_dim.hidden


def valid_output_span(ws, row_idx: int, headers):
    start = to_int(ws.cell(row=row_idx, column=headers[START_COL]).value)
    end = to_int(ws.cell(row=row_idx, column=headers[END_COL]).value)
    if start is None or end is None or start < 0 or end <= start:
        return None
    return start, end


def find_spans(text: str, value: str):
    text_lower = text.casefold()
    value_lower = value.casefold()
    spans = []
    start = 0

    while value_lower:
        pos = text_lower.find(value_lower, start)
        if pos < 0:
            break
        spans.append((pos, pos + len(value)))
        start = pos + 1

    return spans


def choose_text_span(sent_text: str, value: str, output_span=None):
    value = cell_text(value)
    if not sent_text or not value:
        return output_span

    candidates = find_spans(sent_text, value)
    if not candidates:
        return output_span

    if output_span is None:
        return candidates[0]

    output_start, output_end = output_span
    containing = [
        span
        for span in candidates
        if span[0] <= output_start and output_end <= span[1]
    ]
    if containing:
        return sorted(containing, key=lambda span: (span[1] - span[0], span[0]))[0]

    return sorted(
        candidates,
        key=lambda span: (
            abs(span[0] - output_start) + abs(span[1] - output_end),
            span[0],
        ),
    )[0]


def gap_has_connector(gap: str) -> bool:
    return bool(CONNECTOR_GAP_RE.fullmatch(gap))


def gap_can_extend_list(gap: str) -> bool:
    return bool(LIST_EXTENSION_GAP_RE.fullmatch(gap))


def row_group_key(ws, row_idx: int, headers):
    return (
        ws.cell(row=row_idx, column=headers[PMID_COL]).value,
        ws.cell(row=row_idx, column=headers[SENT_INDEX_COL]).value,
        ws.cell(row=row_idx, column=headers[SENT_TEXT_COL]).value,
    )


def build_components(ws, row_indices, headers):
    components = []

    for row_idx in row_indices:
        if DISCARD_COL in headers and is_yes(
            ws.cell(row=row_idx, column=headers[DISCARD_COL]).value
        ):
            continue

        bert_output = cell_text(ws.cell(row=row_idx, column=headers[BERT_OUTPUT_COL]).value)
        bert_base = cell_text(ws.cell(row=row_idx, column=headers[BERT_BASE_COL]).value)
        sent_text = "" if ws.cell(row=row_idx, column=headers[SENT_TEXT_COL]).value is None else str(
            ws.cell(row=row_idx, column=headers[SENT_TEXT_COL]).value
        )
        output_span = valid_output_span(ws, row_idx, headers)
        if not sent_text or not bert_output or not bert_base or output_span is None:
            continue

        base_span = choose_text_span(sent_text, bert_base, output_span)
        if base_span is None:
            base_span = output_span

        components.append(
            {
                "row_idx": row_idx,
                "sent_text": sent_text,
                "bert_output": bert_output,
                "bert_base": bert_base,
                "output_span": output_span,
                "base_span": base_span,
            }
        )

    return sorted(
        components,
        key=lambda item: (
            item["base_span"][0],
            item["base_span"][1],
            item["output_span"][0],
            item["row_idx"],
        ),
    )


def sequence_existing_key(ws, headers, group_key, bert_output, bert_base):
    pmid, sent_index, sent_text = group_key
    return (
        pmid,
        sent_index,
        sent_text,
        normalize_spaces(bert_output).casefold(),
        normalize_spaces(bert_base).casefold(),
    )


def make_existing_sequence_keys(ws, headers):
    keys = set()
    for row_idx in range(2, ws.max_row + 1):
        bert_output = cell_text(ws.cell(row=row_idx, column=headers[BERT_OUTPUT_COL]).value)
        bert_base = cell_text(ws.cell(row=row_idx, column=headers[BERT_BASE_COL]).value)
        if not bert_output or not bert_base:
            continue

        keys.add(sequence_existing_key(ws, headers, row_group_key(ws, row_idx, headers), bert_output, bert_base))

    return keys


def find_merge_sequences(components):
    sequences = []
    if len(components) < 2:
        return sequences

    current = [components[0]]
    current_has_connector = False

    for previous, current_item in zip(components, components[1:]):
        previous_end = previous["base_span"][1]
        current_start = current_item["base_span"][0]
        gap = previous["sent_text"][previous_end:current_start]

        if current_start >= previous_end and (
            gap_has_connector(gap) or gap_can_extend_list(gap)
        ):
            current.append(current_item)
            current_has_connector = current_has_connector or gap_has_connector(gap)
            continue

        if len(current) >= 2 and current_has_connector:
            sequences.append(current)

        current = [current_item]
        current_has_connector = False

    if len(current) >= 2 and current_has_connector:
        sequences.append(current)

    return sequences


def strictly_contains_span(parent_span, child_span) -> bool:
    return (
        parent_span[0] <= child_span[0]
        and parent_span[1] >= child_span[1]
        and parent_span != child_span
    )


def find_nested_numeric_sequences(components):
    sequences = []
    seen_row_sets = set()

    for parent in components:
        contained = [
            component
            for component in components
            if component is not parent
            and strictly_contains_span(parent["base_span"], component["base_span"])
        ]
        if not contained:
            continue

        sequence = sorted(
            [parent] + contained,
            key=lambda item: (
                item["output_span"][0],
                item["output_span"][1],
                item["row_idx"],
            ),
        )
        row_set = frozenset(item["row_idx"] for item in sequence)
        if len(row_set) < 2 or row_set in seen_row_sets:
            continue

        seen_row_sets.add(row_set)
        sequences.append(sequence)

    return sequences


def sequence_values(sequence):
    sent_text = sequence[0]["sent_text"]
    base_start = min(item["base_span"][0] for item in sequence)
    base_end = max(item["base_span"][1] for item in sequence)
    output_start = min(item["output_span"][0] for item in sequence)
    output_end = max(item["output_span"][1] for item in sequence)

    if (
        base_start < 0
        or output_start < 0
        or base_end > len(sent_text)
        or output_end > len(sent_text)
        or base_start >= base_end
        or output_start >= output_end
    ):
        return None

    bert_base = sent_text[base_start:base_end]
    bert_output = sent_text[output_start:output_end]
    return bert_output, bert_base, output_start, output_end


def has_type_or_core_number(text: str) -> bool:
    return bool(TYPE_CORE_NUMBER_RE.search(normalize_spaces(text)))


def starts_with_number(text: str) -> bool:
    return bool(LEADING_NUMBER_RE.search(normalize_spaces(text)))


def component_has_type_core_number(component) -> bool:
    return any(
        has_type_or_core_number(component[field])
        for field in ("bert_base", "bert_output")
    )


def component_starts_with_number(component) -> bool:
    return any(
        starts_with_number(component[field])
        for field in ("bert_base", "bert_output")
    )


def sequence_gap(previous, current_item) -> str:
    previous_base_end = previous["base_span"][1]
    current_base_start = current_item["base_span"][0]

    if current_base_start >= previous_base_end:
        return previous["sent_text"][previous_base_end:current_base_start]

    previous_output_end = previous["output_span"][1]
    current_output_start = current_item["output_span"][0]
    if current_output_start >= previous_output_end:
        return previous["sent_text"][previous_output_end:current_output_start]

    return ""


def sequence_has_type_core_numeric_and_case(sequence) -> bool:
    sequence_has_type_core_number = any(
        component_has_type_core_number(component) for component in sequence
    )

    for previous, current_item in zip(sequence, sequence[1:]):
        gap = sequence_gap(previous, current_item)

        if not re.search(r"\b(?:and|or)\b", gap, flags=re.IGNORECASE):
            continue

        sequence_text = previous["sent_text"][
            min(item["base_span"][0] for item in sequence) :
            max(item["base_span"][1] for item in sequence)
        ]
        if GAP_DIGIT_BEFORE_AND_OR_RE.search(gap) and has_type_or_core_number(sequence_text):
            return True

        previous_is_numeric_continuation = (
            component_has_type_core_number(previous)
            or component_starts_with_number(previous)
        )
        current_is_numeric_continuation = (
            component_has_type_core_number(current_item)
            or component_starts_with_number(current_item)
        )

        if (
            sequence_has_type_core_number
            and previous_is_numeric_continuation
            and current_is_numeric_continuation
        ):
            return True

    return False


def trim_to_type_core_numeric_run(sequence):
    first = None
    for idx, component in enumerate(sequence):
        if component_has_type_core_number(component):
            first = idx
            break

    if first is None:
        return None

    last = first
    for idx in range(first + 1, len(sequence)):
        component = sequence[idx]
        if component_has_type_core_number(component) or component_starts_with_number(component):
            last = idx
            continue
        break

    candidate = sequence[first : last + 1]
    if len(candidate) < 2 or not sequence_has_type_core_numeric_and_case(candidate):
        return None

    return candidate


def process_workbook() -> None:
    if not INPUT_FILE.exists():
        raise FileNotFoundError(f"Input file not found: {INPUT_FILE}")

    wb = load_workbook(INPUT_FILE)
    ws = wb[SHEET_NAME] if SHEET_NAME else wb.active

    headers = get_header_map(ws)
    headers = ensure_column(ws, headers, ADDITION_COL)

    for col_name in REQUIRED_COLUMNS:
        if col_name not in headers:
            raise ValueError(f"Column '{col_name}' not found in workbook.")

    groups = {}
    for row_idx in range(2, ws.max_row + 1):
        groups.setdefault(row_group_key(ws, row_idx, headers), []).append(row_idx)

    existing_sequence_keys = make_existing_sequence_keys(ws, headers)
    row_updates = {}
    rows_to_delete = set()
    claimed_rows = set()

    for group_key, row_indices in groups.items():
        components = build_components(ws, row_indices, headers)
        sequences = find_merge_sequences(components)
        sequences.extend(find_nested_numeric_sequences(components))

        for sequence in sequences:
            sequence = trim_to_type_core_numeric_run(sequence)
            if sequence is None:
                continue

            sequence_rows = {item["row_idx"] for item in sequence}
            if sequence_rows & claimed_rows:
                continue

            values = sequence_values(sequence)
            if values is None:
                continue

            bert_output, bert_base, output_start, output_end = values
            new_key = sequence_existing_key(ws, headers, group_key, bert_output, bert_base)
            if new_key in existing_sequence_keys:
                continue
            existing_sequence_keys.add(new_key)

            keep_row = min(sequence_rows)
            row_updates[keep_row] = {
                "bert_output": bert_output,
                "bert_base": bert_base,
                "start": output_start,
                "end": output_end,
            }
            rows_to_delete.update(row_idx for row_idx in sequence_rows if row_idx != keep_row)
            claimed_rows.update(sequence_rows)

    initial_rows = ws.max_row - 1
    for row_idx, task in row_updates.items():
        ws.cell(row=row_idx, column=headers[BERT_OUTPUT_COL]).value = task["bert_output"]
        ws.cell(row=row_idx, column=headers[BERT_BASE_COL]).value = task["bert_base"]
        ws.cell(row=row_idx, column=headers[ADDITION_COL]).value = "Yes"
        ws.cell(row=row_idx, column=headers[START_COL]).value = task["start"]
        ws.cell(row=row_idx, column=headers[END_COL]).value = task["end"]

    for row_idx in sorted(rows_to_delete, reverse=True):
        ws.delete_rows(row_idx, 1)

    OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT_FILE)

    print(f"Input rows: {initial_rows}")
    print(f"Rows replaced: {len(row_updates)}")
    print(f"Rows removed: {len(rows_to_delete)}")
    print(f"Output rows: {ws.max_row - 1}")
    print(f"Saved output file to: {OUTPUT_FILE}")


if __name__ == "__main__":
    process_workbook()
