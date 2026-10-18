#!/usr/bin/env python3
"""
Remove parenthesized numeric percentages from BERT_Output values, then keep
only pmid, sent_index, sent_text, and BERT_Output in the output workbook.

Examples:
  GM3 (19%) and GM2       -> GM3 and GM2
  value (12.5%), another  -> value, another

Offsets remain 0-based with an exclusive end while the cleanup is performed.
If the cleaned BERT_Output is still a contiguous substring of sent_text,
start_offset/end_offset are updated to that span. If an internal removal makes
the result noncontiguous, both offsets are set to -1. These offset columns,
P_Rem, BERT_Base_NPs, and every other unrequested column are removed before
the final workbook is saved.

Before the final columns are selected, rows with Discard="Yes" are handled by
their original (pmid, sent_index) group size. A discarded row is removed when
that group has another row. A discarded row that is the group's only row is
retained with an empty BERT_Output. BERT_Base_NPs is never changed before it
is removed. Before the final column removal, a BERT_Output beginning with
"hapten" is extended to include an immediately preceding A/B/Y/H/O code,
optionally followed by digits. Sentence-initial A is excluded. Rows with
percentage removal are temporarily marked P_Rem="Yes".
"""

from __future__ import annotations
from glycaner_paths import STAGE3_INPUT, resource, work  # noqa: E402

import re
from copy import copy
from pathlib import Path
from typing import Any
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.utils import get_column_letter


INPUT_FILE = Path(
    work("5_bert_extended.xlsx")
)
OUTPUT_FILE = INPUT_FILE.with_name("6a_percent_removed.xlsx")
SHEET_NAME = None  # Use None to process every worksheet.

PMID_COL = "pmid"
SENT_INDEX_COL = "sent_index"
SENT_TEXT_COL = "sent_text"
BERT_OUTPUT_COL = "BERT_Output"
START_COL = "start_offset"
END_COL = "end_offset"
P_REM_COL = "P_Rem"
DISCARD_COL = "Discard"

FINAL_COLUMNS = [
    PMID_COL,
    SENT_INDEX_COL,
    SENT_TEXT_COL,
    BERT_OUTPUT_COL,
]

# One or more digits, optionally with a decimal part, enclosed in parentheses
# and followed by a percent sign. Optional internal spaces are accepted.
PAREN_PERCENT_RE = re.compile(
    r"\(\s*[+-]?\d+(?:\.\d+)?\s*%\s*\)"
)
SPACE_BEFORE_PUNCT_RE = re.compile(r"\s+([,;:.!?])")
MULTISPACE_RE = re.compile(r"[ \t]{2,}")
HAPTEN_AT_START_RE = re.compile(r"^hapten\b", re.IGNORECASE)
HAPTEN_PREFIX_RE = re.compile(
    r"(?<![A-Za-z0-9])(?P<code>[ABYHO](?:\d+)?)\s+$",
    re.IGNORECASE,
)


def normalize_header(value: Any) -> str:
    return str(value).strip().casefold() if value is not None else ""


def is_readable_xlsx(path: Path) -> bool:
    if not path.is_file():
        return False
    try:
        with ZipFile(path) as archive:
            return archive.testzip() is None
    except (BadZipFile, OSError, TimeoutError):
        return False


def require_readable_input_file(path: Path) -> Path:
    if is_readable_xlsx(path):
        return path
    raise BadZipFile(
        f"Input is not a readable .xlsx file: {path}. If it is stored in "
        "iCloud, choose Download Now in Finder and rerun the script."
    )


def get_headers(ws) -> dict[str, int]:
    return {
        normalize_header(ws.cell(row=1, column=column).value): column
        for column in range(1, ws.max_column + 1)
        if ws.cell(row=1, column=column).value is not None
    }


def copy_cell_format(source: Cell, target: Cell) -> None:
    if source.has_style:
        target.font = copy(source.font)
        target.fill = copy(source.fill)
        target.border = copy(source.border)
        target.alignment = copy(source.alignment)
        target.protection = copy(source.protection)
        target.number_format = source.number_format
    if source.hyperlink:
        target._hyperlink = copy(source.hyperlink)
    if source.comment:
        target.comment = copy(source.comment)


def ensure_column(ws, headers: dict[str, int], column_name: str) -> dict[str, int]:
    key = normalize_header(column_name)
    if key in headers:
        return headers

    column = ws.max_column + 1
    ws.cell(row=1, column=column).value = column_name
    if column > 1:
        previous_column = column - 1
        copy_cell_format(
            ws.cell(row=1, column=previous_column),
            ws.cell(row=1, column=column),
        )
        previous_letter = get_column_letter(previous_column)
        new_letter = get_column_letter(column)
        ws.column_dimensions[new_letter].width = ws.column_dimensions[
            previous_letter
        ].width

        for row_number in range(2, ws.max_row + 1):
            copy_cell_format(
                ws.cell(row=row_number, column=previous_column),
                ws.cell(row=row_number, column=column),
            )

    return get_headers(ws)


def cell_text(value: Any) -> str:
    return "" if value is None else str(value)


def is_yes(value: Any) -> bool:
    return str(value).strip().casefold() == "yes"


def to_int(value: Any) -> int | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    try:
        return int(str(value).strip())
    except (TypeError, ValueError):
        return None


def clean_bert_output(value: str) -> tuple[str, int]:
    matches = list(PAREN_PERCENT_RE.finditer(value))
    if not matches:
        return value, 0

    cleaned = PAREN_PERCENT_RE.sub("", value)
    cleaned = SPACE_BEFORE_PUNCT_RE.sub(r"\1", cleaned)
    cleaned = MULTISPACE_RE.sub(" ", cleaned).strip()
    return cleaned, len(matches)


def find_spans_case_insensitive(text: str, value: str) -> list[tuple[int, int]]:
    if not text or not value:
        return []
    return [
        (match.start(), match.end())
        for match in re.finditer(re.escape(value), text, re.IGNORECASE)
    ]


def choose_cleaned_span(
    sent_text: str,
    cleaned_output: str,
    old_start_value: Any,
    old_end_value: Any,
) -> tuple[int, int] | None:
    spans = find_spans_case_insensitive(sent_text, cleaned_output)
    if not spans:
        return None

    old_start = to_int(old_start_value)
    old_end = to_int(old_end_value)
    if old_start is None or old_end is None:
        return spans[0]

    return min(
        spans,
        key=lambda span: (
            0 if span[0] >= old_start and span[1] <= old_end else 1,
            abs(span[0] - old_start) + abs(span[1] - old_end),
            span[0],
        ),
    )


def apply_discard_rule(ws) -> dict[str, int]:
    """Remove or blank discarded rows using each group's original row count."""
    headers = get_headers(ws)
    pmid_col = headers.get(normalize_header(PMID_COL))
    sent_index_col = headers.get(normalize_header(SENT_INDEX_COL))
    output_col = headers.get(normalize_header(BERT_OUTPUT_COL))
    discard_col = headers.get(normalize_header(DISCARD_COL))

    if None in (pmid_col, sent_index_col, output_col, discard_col):
        return {
            "discarded_rows_removed": 0,
            "singleton_discarded_rows_blanked": 0,
        }

    group_counts: dict[tuple[Any, Any], int] = {}
    for row_number in range(2, ws.max_row + 1):
        key = (
            ws.cell(row=row_number, column=pmid_col).value,
            ws.cell(row=row_number, column=sent_index_col).value,
        )
        group_counts[key] = group_counts.get(key, 0) + 1

    rows_to_remove: list[int] = []
    singleton_rows_to_blank: list[int] = []

    for row_number in range(2, ws.max_row + 1):
        if not is_yes(ws.cell(row=row_number, column=discard_col).value):
            continue

        key = (
            ws.cell(row=row_number, column=pmid_col).value,
            ws.cell(row=row_number, column=sent_index_col).value,
        )
        if group_counts[key] > 1:
            rows_to_remove.append(row_number)
        else:
            singleton_rows_to_blank.append(row_number)

    for row_number in singleton_rows_to_blank:
        ws.cell(row=row_number, column=output_col).value = None

    for row_number in reversed(rows_to_remove):
        ws.delete_rows(row_number)

    return {
        "discarded_rows_removed": len(rows_to_remove),
        "singleton_discarded_rows_blanked": len(singleton_rows_to_blank),
    }


def apply_hapten_prefix_rule(ws) -> int:
    """Extend leading 'hapten' outputs with an adjacent allowed code."""
    headers = get_headers(ws)
    sent_col = headers.get(normalize_header(SENT_TEXT_COL))
    output_col = headers.get(normalize_header(BERT_OUTPUT_COL))
    start_col = headers.get(normalize_header(START_COL))
    end_col = headers.get(normalize_header(END_COL))
    if sent_col is None or output_col is None:
        return 0

    updated_rows = 0
    for row_number in range(2, ws.max_row + 1):
        bert_output = cell_text(
            ws.cell(row=row_number, column=output_col).value
        ).strip()
        if not bert_output or HAPTEN_AT_START_RE.match(bert_output) is None:
            continue

        sent_text = cell_text(ws.cell(row=row_number, column=sent_col).value)
        old_start = (
            ws.cell(row=row_number, column=start_col).value
            if start_col is not None
            else None
        )
        old_end = (
            ws.cell(row=row_number, column=end_col).value
            if end_col is not None
            else None
        )
        output_span = choose_cleaned_span(
            sent_text,
            bert_output,
            old_start,
            old_end,
        )
        if output_span is None:
            continue

        output_start, output_end = output_span
        prefix_match = HAPTEN_PREFIX_RE.search(sent_text[:output_start])
        if prefix_match is None:
            continue

        code = prefix_match.group("code")
        code_start = prefix_match.start("code")
        if (
            code.casefold() == "a"
            and re.search(r"[A-Za-z0-9]", sent_text[:code_start]) is None
        ):
            continue

        ws.cell(row=row_number, column=output_col).value = sent_text[
            code_start:output_end
        ]
        if start_col is not None:
            ws.cell(row=row_number, column=start_col).value = code_start
        if end_col is not None:
            ws.cell(row=row_number, column=end_col).value = output_end
        updated_rows += 1

    return updated_rows


def process_worksheet(ws) -> dict[str, int]:
    headers = get_headers(ws)
    output_col = headers.get(normalize_header(BERT_OUTPUT_COL))
    if output_col is None:
        return {
            "input_rows": ws.max_row - 1,
            "discarded_rows_removed": 0,
            "singleton_discarded_rows_blanked": 0,
            "hapten_prefixes_added": 0,
            "updated_rows": 0,
            "rows_marked": 0,
            "percentage_groups_removed": 0,
            "offsets_updated": 0,
            "offsets_invalidated": 0,
        }

    input_rows = ws.max_row - 1
    discard_summary = apply_discard_rule(ws)
    hapten_prefixes_added = apply_hapten_prefix_rule(ws)
    headers = ensure_column(ws, headers, P_REM_COL)
    output_col = headers[normalize_header(BERT_OUTPUT_COL)]
    sent_col = headers.get(normalize_header(SENT_TEXT_COL))
    start_col = headers.get(normalize_header(START_COL))
    end_col = headers.get(normalize_header(END_COL))
    p_rem_col = headers[normalize_header(P_REM_COL)]

    summary = {
        "input_rows": input_rows,
        **discard_summary,
        "hapten_prefixes_added": hapten_prefixes_added,
        "updated_rows": 0,
        "rows_marked": 0,
        "percentage_groups_removed": 0,
        "offsets_updated": 0,
        "offsets_invalidated": 0,
    }

    for row_number in range(2, ws.max_row + 1):
        original_output = cell_text(
            ws.cell(row=row_number, column=output_col).value
        )
        if not original_output:
            continue

        cleaned_output, removed_count = clean_bert_output(original_output)
        if removed_count == 0:
            continue

        ws.cell(row=row_number, column=output_col).value = cleaned_output or None
        ws.cell(row=row_number, column=p_rem_col).value = "Yes"
        summary["updated_rows"] += 1
        summary["rows_marked"] += 1
        summary["percentage_groups_removed"] += removed_count

        if start_col is None or end_col is None:
            continue

        sent_text = (
            cell_text(ws.cell(row=row_number, column=sent_col).value)
            if sent_col is not None
            else ""
        )
        old_start = ws.cell(row=row_number, column=start_col).value
        old_end = ws.cell(row=row_number, column=end_col).value
        cleaned_span = choose_cleaned_span(
            sent_text,
            cleaned_output,
            old_start,
            old_end,
        )

        if cleaned_span is None:
            ws.cell(row=row_number, column=start_col).value = -1
            ws.cell(row=row_number, column=end_col).value = -1
            summary["offsets_invalidated"] += 1
        else:
            start, end = cleaned_span
            ws.cell(row=row_number, column=start_col).value = start
            ws.cell(row=row_number, column=end_col).value = end
            summary["offsets_updated"] += 1

    return summary


def keep_only_columns(ws, column_names: list[str]) -> int:
    """Delete every worksheet column except column_names."""
    headers = get_headers(ws)
    missing = [
        name for name in column_names if normalize_header(name) not in headers
    ]
    if missing:
        raise ValueError(
            f"Worksheet {ws.title!r} is missing final column(s): "
            f"{', '.join(missing)}"
        )

    keep_keys = {normalize_header(name) for name in column_names}
    columns_to_remove = [
        column
        for column in range(1, ws.max_column + 1)
        if normalize_header(ws.cell(row=1, column=column).value) not in keep_keys
    ]

    for column in reversed(columns_to_remove):
        ws.delete_cols(column)

    final_headers = [
        ws.cell(row=1, column=column).value
        for column in range(1, ws.max_column + 1)
    ]
    if [normalize_header(value) for value in final_headers] != [
        normalize_header(name) for name in column_names
    ]:
        raise ValueError(
            f"Worksheet {ws.title!r} does not have the expected final column "
            f"order: {', '.join(column_names)}"
        )

    if ws.auto_filter and ws.auto_filter.ref:
        ws.auto_filter.ref = ws.dimensions

    return len(columns_to_remove)


def process_workbook(
    input_path: Path,
    output_path: Path,
    sheet_name: str | None = None,
) -> dict[str, int]:
    if not input_path.exists():
        raise FileNotFoundError(f"Input workbook not found: {input_path}")

    workbook = load_workbook(input_path)
    worksheet_names = [sheet_name] if sheet_name else workbook.sheetnames
    totals = {
        "input_rows": 0,
        "discarded_rows_removed": 0,
        "singleton_discarded_rows_blanked": 0,
        "hapten_prefixes_added": 0,
        "updated_rows": 0,
        "rows_marked": 0,
        "percentage_groups_removed": 0,
        "offsets_updated": 0,
        "offsets_invalidated": 0,
        "columns_removed": 0,
    }

    for name in worksheet_names:
        worksheet = workbook[name]
        summary = process_worksheet(worksheet)
        for key, value in summary.items():
            totals[key] += value
        totals["columns_removed"] += keep_only_columns(
            worksheet,
            FINAL_COLUMNS,
        )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return totals


def main() -> None:
    input_path = require_readable_input_file(INPUT_FILE.expanduser().resolve())
    summary = process_workbook(
        input_path,
        OUTPUT_FILE.expanduser().resolve(),
        SHEET_NAME,
    )

    print(f"Input rows: {summary['input_rows']}")
    print(f"Discarded rows removed: {summary['discarded_rows_removed']}")
    print(
        "Singleton discarded rows retained with blank BERT_Output: "
        f"{summary['singleton_discarded_rows_blanked']}"
    )
    print(f"Hapten prefixes added: {summary['hapten_prefixes_added']}")
    print(f"Rows updated: {summary['updated_rows']}")
    print(f"Rows temporarily marked P_Rem=Yes: {summary['rows_marked']}")
    print(
        "Parenthesized percentage groups removed: "
        f"{summary['percentage_groups_removed']}"
    )
    print(f"Offsets updated before column removal: {summary['offsets_updated']}")
    print(f"Offsets set to -1 before column removal: {summary['offsets_invalidated']}")
    print(f"Columns removed: {summary['columns_removed']}")
    print(f"Final columns: {', '.join(FINAL_COLUMNS)}")
    print(f"Saved output file to: {OUTPUT_FILE}")


if __name__ == "__main__":
    main()
