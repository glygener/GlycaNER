#!/usr/bin/env python3
from __future__ import annotations
from glycaner_paths import STAGE3_INPUT, resource, work  # noqa: E402
from collections import OrderedDict
from pathlib import Path
import openpyxl


INFILE = Path(
    work("post_connector_merged.xlsx")
)
OUTFILE = INFILE.with_name("2c_withBaseNPs_post_np.xlsx")

PMID_COL = "pmid"
SENT_INDEX_COL = "sent_index"
SENT_TEXT_COL = "sent_text"
BERT_OUTPUT_COL = "BERT_Output"
START_COL = "start_offset"
END_COL = "end_offset"
BASE_NP_COL = "BERT_Base_NPs"
DISCARD_COL = "Discard"
NP_COL = "NP"

REQUIRED_COLUMNS = {
    PMID_COL,
    SENT_INDEX_COL,
    SENT_TEXT_COL,
    BERT_OUTPUT_COL,
    START_COL,
    END_COL,
    BASE_NP_COL,
}


def is_blank(value) -> bool:
    return value is None or str(value).strip() == ""


def clean(value) -> str:
    return "" if is_blank(value) else str(value).strip()


def is_yes(value) -> bool:
    return str(value).strip().casefold() == "yes"


def to_int(value):
    try:
        return int(value)
    except (TypeError, ValueError):
        return None


def valid_span(row, col_index) -> tuple[int, int] | None:
    if is_blank(row[col_index[BERT_OUTPUT_COL]]):
        return None
    start = to_int(row[col_index[START_COL]])
    end = to_int(row[col_index[END_COL]])
    if start is None or end is None or start < 0 or end <= start:
        return None
    return start, end


def find_base_np_span(sent_text: str, base_np: str) -> tuple[int, int] | None:
    if not sent_text or not base_np:
        return None

    start = sent_text.find(base_np)
    if start >= 0:
        return start, start + len(base_np)

    start = sent_text.casefold().find(base_np.casefold())
    if start >= 0:
        return start, start + len(base_np)
    return None


def group_sentence_rows(rows, col_index):
    grouped = OrderedDict()
    for pos, row in enumerate(rows):
        key = (row[col_index[PMID_COL]], row[col_index[SENT_INDEX_COL]])
        grouped.setdefault(key, []).append((pos, row))
    return grouped


def process_rows(rows, headers):
    col_index = {header: idx for idx, header in enumerate(headers)}

    if NP_COL not in headers:
        headers = headers + [NP_COL]
        for idx, row in enumerate(rows):
            rows[idx] = tuple(list(row) + [""])
        col_index = {header: idx for idx, header in enumerate(headers)}

    output_rows = [list(row) for row in rows]
    remove_positions = set()

    for (_pmid, _sent_index), sentence_rows in group_sentence_rows(rows, col_index).items():
        if DISCARD_COL in col_index and any(
            is_yes(row[col_index[DISCARD_COL]]) for _pos, row in sentence_rows
        ):
            continue

        base_groups = OrderedDict()
        for pos, row in sentence_rows:
            base_np = clean(row[col_index[BASE_NP_COL]])
            bert_output = clean(row[col_index[BERT_OUTPUT_COL]])
            span = valid_span(row, col_index)
            if not base_np or not bert_output or span is None:
                continue
            base_groups.setdefault(base_np, []).append((pos, row, span))

        for base_np, group_rows in base_groups.items():
            distinct_outputs = {
                clean(row[col_index[BERT_OUTPUT_COL]]) for _pos, row, _span in group_rows
            }
            if len(group_rows) < 2 or len(distinct_outputs) < 2:
                continue

            sorted_group = sorted(group_rows, key=lambda item: (item[2][0], item[2][1]))
            first_pos, first_row, first_span = sorted_group[0]
            _last_pos, last_row, last_span = sorted_group[-1]

            sent_value = first_row[col_index[SENT_TEXT_COL]]
            sent_text = "" if sent_value is None else str(sent_value)
            if not sent_text:
                continue

            base_np_span = find_base_np_span(sent_text, base_np)
            if base_np_span is not None:
                base_start, base_end = base_np_span
                base_contained_rows = []
                for pos, row in sentence_rows:
                    span = valid_span(row, col_index)
                    if span is None:
                        continue
                    row_start, row_end = span
                    if row_start >= base_start and row_end <= base_end:
                        base_contained_rows.append((pos, row_start, row_end))

                if base_contained_rows:
                    start = min(row_start for _pos, row_start, _row_end in base_contained_rows)
                    end = max(row_end for _pos, _row_start, row_end in base_contained_rows)
                else:
                    start = first_span[0]
                    end = last_span[1]
            else:
                start = first_span[0]
                end = last_span[1]

            if start < 0 or end > len(sent_text) or start >= end:
                continue

            merged_text = sent_text[start:end]
            contained_rows = []
            for pos, row in sentence_rows:
                span = valid_span(row, col_index)
                if span is None:
                    continue
                row_start, row_end = span
                if row_start >= start and row_end <= end:
                    contained_rows.append((pos, row_start, row_end))

            if not contained_rows:
                contained_rows = [
                    (pos, span[0], span[1]) for pos, _row, span in sorted_group
                ]

            keep_pos, _keep_start, _keep_end = sorted(
                contained_rows, key=lambda item: (item[1], item[2])
            )[0]

            output_rows[keep_pos][col_index[BERT_OUTPUT_COL]] = merged_text
            output_rows[keep_pos][col_index[START_COL]] = start
            output_rows[keep_pos][col_index[END_COL]] = end
            output_rows[keep_pos][col_index[BASE_NP_COL]] = base_np
            output_rows[keep_pos][col_index[NP_COL]] = "Yes"

            for pos, _row_start, _row_end in contained_rows:
                if pos != keep_pos:
                    remove_positions.add(pos)

    kept_rows = [
        row for pos, row in enumerate(output_rows) if pos not in remove_positions
    ]
    return headers, kept_rows


def main() -> None:
    wb = openpyxl.load_workbook(INFILE, data_only=False)
    out_wb = openpyxl.Workbook()
    out_wb.remove(out_wb.active)

    for ws in wb.worksheets:
        out_ws = out_wb.create_sheet(ws.title)
        headers = [cell.value for cell in ws[1]]
        rows = list(ws.iter_rows(min_row=2, values_only=True))

        if REQUIRED_COLUMNS.issubset(set(headers)):
            headers, rows = process_rows(rows, headers)

        out_ws.append(headers)
        for row in rows:
            out_ws.append(row)

    OUTFILE.parent.mkdir(parents=True, exist_ok=True)
    out_wb.save(OUTFILE)
    print(f"Saved: {OUTFILE}")


if __name__ == "__main__":
    main()
