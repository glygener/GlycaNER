#!/usr/bin/env python3
"""
Extend BERT_Output using regex features found only in BERT_Base_NPs.

For every non-discarded row:
1. Locate BERT_Output inside BERT_Base_NPs.
2. Search the BaseNP text to the left and right of BERT_Output with every
   regex stored in Dict_Normalization/*.txt.
3. Extend a side only when that side has at least one match from a file other
   than bi_di_tri.txt or digit_modifier.txt. Matches from those two files may
   participate when another allowed feature also matches the same side. The
   para, iso, lipo, and pauci rules from unknown.txt behave the same way.
4. A left extension begins at the full word containing its earliest matched
   feature, so partial results such as "parate" or "binant" are not produced.
   A right extension ends at the full word containing its rightmost match, so
   unmatched tails such as "animals", "species", or "product" are not added.
5. Ignore the standalone cer/cer(s) headword regex because it creates false
   matches inside words such as "glyceryl" and "cancer".
6. On the right, also extend through alpha/beta when it follows a digit.
7. Search no more than the three nearest words on either side for ordinary
   dictionary matches. Alpha/beta, GlcNAc, GalNAc, NeuAc, NeuGc, Neu5Ac,
   Neu5Gc, ManNAc, Gal, Glc, sialyl, and a leading isotope form such as
   [3H] or [9-3H] may extend beyond that window on either side. Alpha/beta
   sequences ending in Gn-T5 or GnT are excluded from extension.
8. On the right, also include alpha/beta followed by numbers, Gal alpha/beta
   followed by numbers, and/or followed by numbers, and standalone uppercase
   A/B/O/H/I letters optionally preceded by and/or.
9. Short chemical/modifier patterns such as Ac and NH must match complete
   tokens; they cannot match inside words such as activity or enhancement.
10. Complete a word cut at the existing output boundary, preserve an existing
    output prefix immediately before an overlapping BaseNP, and include a
    terminal linkage letter such as the R in beta 1-3-R.
11. Include an adjacent opening '(' or '[' and closing ')' or ']'.
12. Store the exact added substrings in new Left and Right columns and update
   BERT_Output with 0-based, end-exclusive sentence offsets.

All worksheets and rows are retained. Rows with Discard="Yes" are untouched.
"""

from __future__ import annotations
from glycaner_paths import STAGE3_INPUT, resource, work  # noqa: E402

import re
from copy import copy
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.utils import get_column_letter


INPUT_FILE = Path(
    work("4_merging_updated.xlsx")
)
OUTPUT_FILE = INPUT_FILE.with_name("5_bert_extended.xlsx")

DICT_FOLDER = Path(
    resource("Dict_Normalization")
)
SHEET_NAME = None  # Use None to process every worksheet with required columns.

SENT_TEXT_COL = "sent_text"
BERT_OUTPUT_COL = "BERT_Output"
BERT_BASE_COL = "BERT_Base_NPs"
START_COL = "start_offset"
END_COL = "end_offset"
DISCARD_COL = "Discard"
LEFT_COL = "Left"
RIGHT_COL = "Right"

REQUIRED_COLUMNS = {SENT_TEXT_COL, BERT_OUTPUT_COL, BERT_BASE_COL}
NON_EXTENDING_FILES = {"bi_di_tri.txt", "digit_modifier.txt"}
NON_EXTENDING_RULES = {
    ("unknown.txt", "(?i)para"),
    ("unknown.txt", "(?i)iso"),
    ("unknown.txt", "(?i)lipo"),
    ("unknown.txt", "(?i)pauci"),
    ("globo.txt", "(?i)neo"),
    ("globo.txt", "(?i)cis"),
    ("globo.txt", "(?i)lyso"),
}
IGNORED_RULES = {("headwords.txt", "(?i)cer(s)?")}
WHOLE_TOKEN_RULES = {
    ("digit_modifier.txt", r"\b([Aa]c|[Aa]s|[Ss]u|[Cc]is)-?"),
    ("chemical.txt", r"(?i)-?-?-?NH(4)?-?"),
    ("chemical.txt", r"(?i)-?-?-?OH(4)?-?"),
    ("chemical.txt", r"(?i)-?-?-?COO-?"),
    ("chemical.txt", r"(?i)-?-?-?COOH-?"),
}
CONDITIONAL_MATCH_TEXT_RE = re.compile(
    r"(?:para|iso|lipo|pauci)",
    re.IGNORECASE,
)
LINKAGE_NUMBER_PATTERN = (
    r"\d+[A-Za-z]*"
    r"(?:\s*[,./'\u2032-]\s*\d+[A-Za-z]*)*"
    r"(?:-\s*[A-Za-z]+\d+)*"
)
RIGHT_DIGIT_ALPHA_BETA_RE = re.compile(
    r"\d[^A-Za-z0-9]{0,20}(?:alpha|beta|\u03b1|\u03b2)(?![A-Za-z])",
    re.IGNORECASE,
)
RIGHT_ALPHA_BETA_NUMBER_RE = re.compile(
    r"(?:alpha|beta|\u03b1|\u03b2)"
    r"\s*[-(]?\s*" + LINKAGE_NUMBER_PATTERN + r"(?:\s*\))?",
    re.IGNORECASE,
)
RIGHT_CONJUNCTION_NUMBER_RE = re.compile(
    r"\b(?:and|or)\b"
    r"\s*[-(]?\s*" + LINKAGE_NUMBER_PATTERN + r"(?:\s*\))?",
    re.IGNORECASE,
)
RIGHT_STANDALONE_LETTER_RE = re.compile(
    r"(?:(?:\band\b|\bor\b)\s*)?"
    r"\b[ABOHI](?:\s*/\s*[ABOHI])*\b"
)
BIOCHEMICAL_EXCEPTION_RE = re.compile(
    r"(?<![A-Za-z])(?:"
    r"GlcNAc|GalNAc|Neu5Ac|Neu5Gc|NeuAc|NeuGc|ManNAc|"
    r"Gal|Glc|sialyl|"
    r"(?:alpha|beta|\u03b1|\u03b2)"
    r"(?:\s*[-(]?\s*" + LINKAGE_NUMBER_PATTERN + r"(?:\s*\))?)?"
    r")(?![A-Za-z])",
    re.IGNORECASE,
)
LEADING_NUMBER_H_RE = re.compile(
    r"^\s*\[\s*\d+(?:\s*-\s*\d+)*\s*H\s*\]",
    re.IGNORECASE,
)
ALPHA_BETA_NUMBER_AT_END_RE = re.compile(
    r"(?:alpha|beta|\u03b1|\u03b2)"
    r"\s*[-(]?\s*" + LINKAGE_NUMBER_PATTERN + r"(?:\s*\))?\s*$",
    re.IGNORECASE,
)
TERMINAL_LINKAGE_LETTER_RE = re.compile(
    r"^\s*-\s*[A-Za-z](?![A-Za-z0-9])"
)
EXCLUDED_GNT_SEQUENCE_RE = re.compile(
    r"(?<![A-Za-z])(?:alpha|beta|\u03b1|\u03b2)"
    r"\s*[-(]?\s*(?:\d+\s*)?Gn(?:-T5|T)\b",
    re.IGNORECASE,
)
WORD_RE = re.compile(
    r"[A-Za-z0-9\u03b1-\u03c9\u0391-\u03a9]+"
    r"(?:[-'\u2032][A-Za-z0-9\u03b1-\u03c9\u0391-\u03a9]+)*"
)
SIDE_WORD_LIMIT = 3


@dataclass(frozen=True)
class RegexRule:
    source_file: str
    line_number: int
    pattern_text: str
    regex: re.Pattern
    non_extending: bool
    whole_token: bool


@dataclass(frozen=True)
class SideMatch:
    source_file: str
    start: int
    end: int
    non_extending: bool


@dataclass(frozen=True)
class EntityLocation:
    output_rel_start: int
    output_rel_end: int
    base_sentence_start: int | None
    output_sentence_span: tuple[int, int] | None


@dataclass(frozen=True)
class RowUpdate:
    bert_output: str
    left_added: str | None
    right_added: str | None
    start: int | None
    end: int | None


def normalize_header(value: Any) -> str:
    return str(value).strip().casefold() if value is not None else ""


def clean_cell(value: Any) -> str:
    return "" if value is None else str(value).strip()


def sentence_text(value: Any) -> str:
    return "" if value is None else str(value)


def is_yes(value: Any) -> bool:
    return str(value).strip().casefold() == "yes"


def to_int(value: Any) -> int | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    try:
        return int(str(value).strip())
    except (TypeError, ValueError):
        return None


def copy_cell_format(source: Cell, target: Cell) -> None:
    if source.has_style:
        target.font = copy(source.font)
        target.fill = copy(source.fill)
        target.border = copy(source.border)
        target.alignment = copy(source.alignment)
        target.protection = copy(source.protection)
        target.number_format = source.number_format
    if source.hyperlink:
        target._hyperlink = copy(source.hyperlink)
    if source.comment:
        target.comment = copy(source.comment)


def get_headers(ws) -> dict[str, int]:
    return {
        normalize_header(ws.cell(row=1, column=column).value): column
        for column in range(1, ws.max_column + 1)
        if ws.cell(row=1, column=column).value is not None
    }


def ensure_column(ws, headers: dict[str, int], column_name: str) -> dict[str, int]:
    key = normalize_header(column_name)
    if key in headers:
        return headers

    column = ws.max_column + 1
    ws.cell(row=1, column=column).value = column_name

    if column > 1:
        previous_column = column - 1
        copy_cell_format(
            ws.cell(row=1, column=previous_column),
            ws.cell(row=1, column=column),
        )
        previous_letter = get_column_letter(previous_column)
        new_letter = get_column_letter(column)
        ws.column_dimensions[new_letter].width = ws.column_dimensions[
            previous_letter
        ].width

        for row_number in range(2, ws.max_row + 1):
            copy_cell_format(
                ws.cell(row=row_number, column=previous_column),
                ws.cell(row=row_number, column=column),
            )

    return get_headers(ws)


def load_regex_rules(folder: Path) -> list[RegexRule]:
    if not folder.is_dir():
        raise FileNotFoundError(f"Dictionary folder not found: {folder}")

    text_files = sorted(folder.glob("*.txt"))
    if not text_files:
        raise FileNotFoundError(f"No .txt dictionary files found in: {folder}")

    rules: list[RegexRule] = []
    for text_file in text_files:
        lines = text_file.read_text(encoding="utf-8-sig").splitlines()
        for line_number, line in enumerate(lines, start=1):
            pattern_text = line.strip()
            if not pattern_text or pattern_text.startswith("#"):
                continue
            if (text_file.name, pattern_text) in IGNORED_RULES:
                continue

            try:
                regex = re.compile(pattern_text)
            except re.error as error:
                raise ValueError(
                    f"Invalid regex in {text_file.name}, line {line_number}: "
                    f"{pattern_text!r}: {error}"
                ) from error

            rules.append(
                RegexRule(
                    source_file=text_file.name,
                    line_number=line_number,
                    pattern_text=pattern_text,
                    regex=regex,
                    non_extending=(
                        text_file.name in NON_EXTENDING_FILES
                        or (text_file.name, pattern_text) in NON_EXTENDING_RULES
                    ),
                    whole_token=(
                        (text_file.name, pattern_text) in WHOLE_TOKEN_RULES
                    ),
                )
            )

    if not rules:
        raise ValueError(f"No usable regex patterns found in: {folder}")
    return rules


def find_case_insensitive_spans(text: str, value: str) -> list[tuple[int, int]]:
    if not text or not value:
        return []
    return [
        (match.start(), match.end())
        for match in re.finditer(re.escape(value), text, re.IGNORECASE)
    ]


def valid_sentence_span(
    sent_text: str,
    start_value: Any,
    end_value: Any,
) -> tuple[int, int] | None:
    start = to_int(start_value)
    end = to_int(end_value)
    if start is None or end is None or start < 0 or end <= start:
        return None
    if end > len(sent_text):
        return None
    return start, end


def span_distance(first: tuple[int, int], second: tuple[int, int]) -> int:
    return abs(first[0] - second[0]) + abs(first[1] - second[1])


def choose_base_sentence_span(
    sent_text: str,
    bert_base: str,
    output_span: tuple[int, int] | None,
) -> tuple[int, int] | None:
    base_spans = find_case_insensitive_spans(sent_text, bert_base)
    if not base_spans:
        return None
    if output_span is None:
        return base_spans[0]

    containing = [
        span
        for span in base_spans
        if span[0] <= output_span[0] and span[1] >= output_span[1]
    ]
    if containing:
        return min(containing, key=lambda span: (span[1] - span[0], span[0]))
    return min(base_spans, key=lambda span: span_distance(span, output_span))


def locate_output_in_base(
    sent_text: str,
    bert_output: str,
    bert_base: str,
    start_value: Any,
    end_value: Any,
) -> EntityLocation | None:
    relative_spans = find_case_insensitive_spans(bert_base, bert_output)
    if not relative_spans:
        return None

    output_span = valid_sentence_span(sent_text, start_value, end_value)
    base_span = choose_base_sentence_span(sent_text, bert_base, output_span)

    if base_span is not None and output_span is not None:
        target_relative_start = output_span[0] - base_span[0]
        relative_span = min(
            relative_spans,
            key=lambda span: (
                abs(span[0] - target_relative_start),
                span[0],
            ),
        )
    else:
        relative_span = relative_spans[0]

    if base_span is None and output_span is not None:
        possible_base_start = output_span[0] - relative_span[0]
        possible_base_end = possible_base_start + len(bert_base)
        if (
            0 <= possible_base_start < possible_base_end <= len(sent_text)
            and sent_text[possible_base_start:possible_base_end].casefold()
            == bert_base.casefold()
        ):
            base_span = (possible_base_start, possible_base_end)

    return EntityLocation(
        output_rel_start=relative_span[0],
        output_rel_end=relative_span[1],
        base_sentence_start=base_span[0] if base_span is not None else None,
        output_sentence_span=output_span,
    )


def extension_base_text(
    sent_text: str,
    bert_output: str,
    bert_base: str,
    start_value: Any,
    end_value: Any,
) -> str:
    """Preserve an output prefix immediately before an overlapping BaseNP."""
    if find_case_insensitive_spans(bert_base, bert_output):
        return bert_base

    output_folded = bert_output.casefold()
    base_folded = bert_base.casefold()
    maximum_overlap = min(len(bert_output), len(bert_base))
    output_span = valid_sentence_span(sent_text, start_value, end_value)

    for overlap_length in range(maximum_overlap, 3, -1):
        if output_folded[-overlap_length:] != base_folded[:overlap_length]:
            continue

        candidate = bert_output[:-overlap_length] + bert_base
        if output_span is not None:
            candidate_end = output_span[0] + len(candidate)
            if (
                candidate_end <= len(sent_text)
                and sent_text[output_span[0]:candidate_end].casefold()
                == candidate.casefold()
            ):
                return candidate
        elif find_case_insensitive_spans(sent_text, candidate):
            return candidate

    return bert_base


def token_core_span(text: str, start: int, end: int) -> tuple[int, int]:
    while start < end and not text[start].isalnum():
        start += 1
    while end > start and not text[end - 1].isalnum():
        end -= 1
    return start, end


def has_complete_token_boundaries(text: str, start: int, end: int) -> bool:
    core_start, core_end = token_core_span(text, start, end)
    if core_start >= core_end:
        return False
    left_ok = core_start == 0 or not text[core_start - 1].isalnum()
    right_ok = core_end == len(text) or not text[core_end].isalnum()
    return left_ok and right_ok


def starts_at_token_boundary(text: str, start: int, end: int) -> bool:
    core_start, core_end = token_core_span(text, start, end)
    if core_start >= core_end:
        return False
    return core_start == 0 or not text[core_start - 1].isalnum()


def side_search_bounds(
    side_text: str,
    is_left: bool,
    apply_word_limit: bool,
) -> tuple[int, int]:
    if not apply_word_limit:
        return 0, len(side_text)

    words = list(WORD_RE.finditer(side_text))
    if len(words) <= SIDE_WORD_LIMIT:
        return 0, len(side_text)
    if is_left:
        return words[-SIDE_WORD_LIMIT].start(), len(side_text)
    return 0, words[SIDE_WORD_LIMIT - 1].end()


def find_side_matches(
    side_text: str,
    rules: list[RegexRule],
    search_start: int = 0,
    search_end: int | None = None,
) -> list[SideMatch]:
    if not side_text.strip():
        return []

    if search_end is None:
        search_end = len(side_text)
    search_text = side_text[search_start:search_end]
    excluded_spans = [
        (match.start(), match.end())
        for match in EXCLUDED_GNT_SEQUENCE_RE.finditer(side_text)
    ]

    matches: list[SideMatch] = []
    for rule in rules:
        for match in rule.regex.finditer(search_text):
            if match.start() == match.end():
                continue
            match_start = search_start + match.start()
            match_end = search_start + match.end()
            if any(
                match_start < excluded_end and excluded_start < match_end
                for excluded_start, excluded_end in excluded_spans
            ):
                continue
            if rule.whole_token and not has_complete_token_boundaries(
                side_text,
                match_start,
                match_end,
            ):
                continue
            non_extending = (
                rule.non_extending
                or CONDITIONAL_MATCH_TEXT_RE.fullmatch(match.group())
                is not None
            )
            if non_extending and not starts_at_token_boundary(
                side_text,
                match_start,
                match_end,
            ):
                continue
            matches.append(
                SideMatch(
                    source_file=rule.source_file,
                    start=match_start,
                    end=match_end,
                    non_extending=non_extending,
                )
            )
    return matches


def find_biochemical_exception_matches(side_text: str) -> list[SideMatch]:
    excluded_spans = [
        (match.start(), match.end())
        for match in EXCLUDED_GNT_SEQUENCE_RE.finditer(side_text)
    ]
    matches = []
    for match in BIOCHEMICAL_EXCEPTION_RE.finditer(side_text):
        if any(
            match.start() < excluded_end and excluded_start < match.end()
            for excluded_start, excluded_end in excluded_spans
        ):
            continue
        matches.append(
            SideMatch(
                source_file="__biochemical_exception__",
                start=match.start(),
                end=match.end(),
                non_extending=False,
            )
        )

    leading_number_h = LEADING_NUMBER_H_RE.search(side_text)
    if leading_number_h is not None:
        matches.append(
            SideMatch(
                source_file="__leading_number_h__",
                start=leading_number_h.start(),
                end=leading_number_h.end(),
                non_extending=False,
            )
        )
    return matches


def side_can_extend(matches: list[SideMatch]) -> bool:
    return any(not match.non_extending for match in matches)


def full_word_start(text: str, start: int) -> int:
    while start > 0 and (text[start - 1].isalnum() or text[start - 1] == "_"):
        start -= 1
    return start


def full_word_end(text: str, end: int) -> int:
    while end < len(text) and (text[end].isalnum() or text[end] == "_"):
        end += 1
    return end


def split_token_extension_end(text: str, output_end: int) -> int | None:
    """Finish a biochemical token cut by the current BERT_Output edge."""
    if not (0 < output_end < len(text)):
        return None
    if not (text[output_end - 1].isalnum() and text[output_end].isalnum()):
        return None

    end = full_word_end(text, output_end)
    while end < len(text) and text[end] == "-":
        suffix_start = end + 1
        suffix_end = full_word_end(text, suffix_start)
        if suffix_end == suffix_start:
            break
        end = suffix_end
    return end


def include_terminal_linkage_letter(text: str, end: int) -> int:
    if not ALPHA_BETA_NUMBER_AT_END_RE.search(text[:end]):
        return end
    match = TERMINAL_LINKAGE_LETTER_RE.match(text[end:])
    return end + match.end() if match is not None else end


def include_opening_brackets(text: str, start: int) -> int:
    while start > 0:
        before_whitespace = start
        while before_whitespace > 0 and text[before_whitespace - 1].isspace():
            before_whitespace -= 1
        if before_whitespace > 0 and text[before_whitespace - 1] in "([":
            start = before_whitespace - 1
            continue
        break
    return start


def include_closing_brackets(text: str, end: int) -> int:
    while end < len(text):
        after_whitespace = end
        while after_whitespace < len(text) and text[after_whitespace].isspace():
            after_whitespace += 1
        if after_whitespace < len(text) and text[after_whitespace] in ")]":
            end = after_whitespace + 1
            continue
        break
    return end


def right_special_extension_end(
    right_text: str,
    search_end: int,
) -> int | None:
    search_text = right_text[:search_end]
    excluded_spans = [
        (match.start(), match.end())
        for match in EXCLUDED_GNT_SEQUENCE_RE.finditer(right_text)
    ]
    ends = [
        match.end()
        for regex in (
            RIGHT_DIGIT_ALPHA_BETA_RE,
            RIGHT_ALPHA_BETA_NUMBER_RE,
            RIGHT_CONJUNCTION_NUMBER_RE,
            RIGHT_STANDALONE_LETTER_RE,
        )
        for match in regex.finditer(search_text)
        if not any(
            match.start() < excluded_end and excluded_start < match.end()
            for excluded_start, excluded_end in excluded_spans
        )
    ]
    return max(ends, default=None)


def trim_relative_span(text: str, start: int, end: int) -> tuple[int, int] | None:
    while start < end and text[start].isspace():
        start += 1
    while end > start and text[end - 1].isspace():
        end -= 1
    return (start, end) if start < end else None


def build_row_update(
    sent_text: str,
    bert_output: str,
    bert_base: str,
    start_value: Any,
    end_value: Any,
    rules: list[RegexRule],
) -> tuple[RowUpdate | None, bool]:
    """Return (update, span_failed) for one row."""
    extension_base = extension_base_text(
        sent_text,
        bert_output,
        bert_base,
        start_value,
        end_value,
    )
    location = locate_output_in_base(
        sent_text,
        bert_output,
        extension_base,
        start_value,
        end_value,
    )
    if location is None:
        return None, False

    left_text = extension_base[: location.output_rel_start]
    right_text = extension_base[location.output_rel_end :]

    apply_word_limit = True
    left_search_start, left_search_end = side_search_bounds(
        left_text,
        is_left=True,
        apply_word_limit=apply_word_limit,
    )
    right_search_start, right_search_end = side_search_bounds(
        right_text,
        is_left=False,
        apply_word_limit=apply_word_limit,
    )

    left_matches = find_side_matches(
        left_text,
        rules,
        left_search_start,
        left_search_end,
    )
    right_matches = find_side_matches(
        right_text,
        rules,
        right_search_start,
        right_search_end,
    )
    left_matches.extend(find_biochemical_exception_matches(left_text))
    right_matches.extend(find_biochemical_exception_matches(right_text))
    extend_left = side_can_extend(left_matches)
    extend_right = side_can_extend(right_matches)
    special_right_end = right_special_extension_end(
        right_text,
        right_search_end,
    )
    split_token_end = split_token_extension_end(
        extension_base,
        location.output_rel_end,
    )
    if (
        not extend_left
        and not extend_right
        and special_right_end is None
        and split_token_end is None
    ):
        return None, False

    new_relative_start = (
        full_word_start(
            extension_base,
            min(match.start for match in left_matches),
        )
        if extend_left
        else location.output_rel_start
    )
    right_end_candidates = [location.output_rel_end]
    if extend_right:
        rightmost_match_end = max(match.end for match in right_matches)
        right_end_candidates.append(
            location.output_rel_end
            + full_word_end(right_text, rightmost_match_end)
        )
    if special_right_end is not None:
        right_end_candidates.append(
            location.output_rel_end
            + full_word_end(right_text, special_right_end)
        )
    if split_token_end is not None:
        right_end_candidates.append(split_token_end)
    new_relative_end = max(right_end_candidates)
    new_relative_end = include_terminal_linkage_letter(
        extension_base,
        new_relative_end,
    )

    new_relative_start = include_opening_brackets(extension_base, new_relative_start)
    new_relative_end = include_closing_brackets(extension_base, new_relative_end)
    trimmed = trim_relative_span(
        extension_base,
        new_relative_start,
        new_relative_end,
    )
    if trimmed is None:
        return None, False
    new_relative_start, new_relative_end = trimmed

    expected_output = extension_base[new_relative_start:new_relative_end]
    new_start: int | None = None
    new_end: int | None = None

    if location.base_sentence_start is not None:
        new_start = location.base_sentence_start + new_relative_start
        new_end = location.base_sentence_start + new_relative_end
    elif location.output_sentence_span is not None:
        output_start, _ = location.output_sentence_span
        candidate_base_start = output_start - location.output_rel_start
        new_start = candidate_base_start + new_relative_start
        new_end = candidate_base_start + new_relative_end

    if new_start is None or new_end is None:
        return None, True
    if not (0 <= new_start < new_end <= len(sent_text)):
        return None, True

    sentence_output = sent_text[new_start:new_end]
    if sentence_output.casefold() != expected_output.casefold():
        return None, True

    left_added = (
        extension_base[new_relative_start : location.output_rel_start].strip()
        if new_relative_start < location.output_rel_start
        else None
    )
    right_added = (
        extension_base[location.output_rel_end : new_relative_end].strip()
        if new_relative_end > location.output_rel_end
        else None
    )
    return (
        RowUpdate(
            bert_output=sentence_output,
            left_added=left_added or None,
            right_added=right_added or None,
            start=new_start,
            end=new_end,
        ),
        False,
    )


def process_worksheet(ws, rules: list[RegexRule]) -> dict[str, int]:
    headers = get_headers(ws)
    required_keys = {normalize_header(column) for column in REQUIRED_COLUMNS}
    if not required_keys.issubset(headers):
        return {
            "input_rows": ws.max_row - 1,
            "updated_rows": 0,
            "left_extensions": 0,
            "right_extensions": 0,
            "both_extensions": 0,
            "discarded_rows_skipped": 0,
            "span_failures": 0,
        }

    headers = ensure_column(ws, headers, LEFT_COL)
    headers = ensure_column(ws, headers, RIGHT_COL)

    sent_col = headers[normalize_header(SENT_TEXT_COL)]
    output_col = headers[normalize_header(BERT_OUTPUT_COL)]
    base_col = headers[normalize_header(BERT_BASE_COL)]
    start_col = headers.get(normalize_header(START_COL))
    end_col = headers.get(normalize_header(END_COL))
    discard_col = headers.get(normalize_header(DISCARD_COL))
    left_col = headers[normalize_header(LEFT_COL)]
    right_col = headers[normalize_header(RIGHT_COL)]

    summary = {
        "input_rows": ws.max_row - 1,
        "updated_rows": 0,
        "left_extensions": 0,
        "right_extensions": 0,
        "both_extensions": 0,
        "discarded_rows_skipped": 0,
        "span_failures": 0,
    }

    for row_number in range(2, ws.max_row + 1):
        if discard_col is not None and is_yes(
            ws.cell(row=row_number, column=discard_col).value
        ):
            summary["discarded_rows_skipped"] += 1
            continue

        sent_text = sentence_text(ws.cell(row=row_number, column=sent_col).value)
        bert_output = clean_cell(ws.cell(row=row_number, column=output_col).value)
        bert_base = clean_cell(ws.cell(row=row_number, column=base_col).value)
        if not sent_text.strip() or not bert_output or not bert_base:
            continue

        start_value = (
            ws.cell(row=row_number, column=start_col).value
            if start_col is not None
            else None
        )
        end_value = (
            ws.cell(row=row_number, column=end_col).value
            if end_col is not None
            else None
        )

        update, span_failed = build_row_update(
            sent_text,
            bert_output,
            bert_base,
            start_value,
            end_value,
            rules,
        )
        if span_failed:
            summary["span_failures"] += 1
        if update is None:
            continue

        ws.cell(row=row_number, column=output_col).value = update.bert_output
        ws.cell(row=row_number, column=left_col).value = update.left_added
        ws.cell(row=row_number, column=right_col).value = update.right_added
        if start_col is not None:
            ws.cell(row=row_number, column=start_col).value = update.start
        if end_col is not None:
            ws.cell(row=row_number, column=end_col).value = update.end

        summary["updated_rows"] += 1
        if update.left_added is not None:
            summary["left_extensions"] += 1
        if update.right_added is not None:
            summary["right_extensions"] += 1
        if update.left_added is not None and update.right_added is not None:
            summary["both_extensions"] += 1

    if ws.auto_filter and ws.auto_filter.ref:
        ws.auto_filter.ref = ws.dimensions

    return summary


def process_workbook(
    input_path: Path,
    output_path: Path,
    dictionary_folder: Path,
    sheet_name: str | None = None,
) -> dict[str, int]:
    if not input_path.exists():
        raise FileNotFoundError(f"Input workbook not found: {input_path}")

    rules = load_regex_rules(dictionary_folder)
    workbook = load_workbook(input_path)
    worksheet_names = [sheet_name] if sheet_name else workbook.sheetnames

    totals = {
        "dictionary_files": len({rule.source_file for rule in rules}),
        "dictionary_patterns": len(rules),
        "input_rows": 0,
        "updated_rows": 0,
        "left_extensions": 0,
        "right_extensions": 0,
        "both_extensions": 0,
        "discarded_rows_skipped": 0,
        "span_failures": 0,
    }

    for name in worksheet_names:
        summary = process_worksheet(workbook[name], rules)
        for key, value in summary.items():
            totals[key] += value

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return totals


def main() -> None:
    summary = process_workbook(
        INPUT_FILE.expanduser().resolve(),
        OUTPUT_FILE.expanduser().resolve(),
        DICT_FOLDER.expanduser().resolve(),
        SHEET_NAME,
    )

    print(f"Dictionary files: {summary['dictionary_files']}")
    print(f"Regex patterns: {summary['dictionary_patterns']}")
    print(f"Input rows: {summary['input_rows']}")
    print(f"Rows updated: {summary['updated_rows']}")
    print(f"Rows extended left: {summary['left_extensions']}")
    print(f"Rows extended right: {summary['right_extensions']}")
    print(f"Rows extended on both sides: {summary['both_extensions']}")
    print(f"Discarded rows skipped: {summary['discarded_rows_skipped']}")
    print(f"Eligible rows skipped for span failures: {summary['span_failures']}")
    print(f"Saved output file to: {OUTPUT_FILE}")


if __name__ == "__main__":
    main()
