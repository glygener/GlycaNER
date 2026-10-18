#!/usr/bin/env python3
from glycaner_paths import STAGE3_INPUT, resource, work  # noqa: E402
import re
from copy import copy
from pathlib import Path

from openpyxl import load_workbook


# ---------- HARD-CODED FILE PATHS ----------
INPUT_FILE = Path(
    work("3_withMergedBaseNPs.xlsx")
)
OUTPUT_FILE = Path(
    work("4_merging_updated.xlsx")
)
SHEET_NAME = None


# ---------- configurable names ----------
BERT_OUTPUT_COL = "BERT_Output"
BERT_BASE_COL = "BERT_Base_NPs"
MERGED_NP_COL = "Merged_Base_NPs"
MERGE_FLAG_COL = "Merge"

START_COL = "start_offset"
END_COL = "end_offset"

PMID_COL = "pmid"
SENT_INDEX_COL = "sent_index"
SENT_TEXT_COL = "sent_text"
# ----------------------------------------


OF_ALLOWED = {"of", "of a", "of an", "of the"}

WITH_ALLOWED = {
    "with",
    "with a",
    "with an",
    "with the",
}

IN_ON_ALLOWED = {
    "in",
    "in a",
    "in an",
    "in the",
    "on",
    "on a",
    "on an",
    "on the",
}

STRUCTURE_WORDS = {"structure", "structures"}

ALLOWED_STRUCTURE_TERM_RE = re.compile(
    r"\b(?:glycans?|[A-Za-z]*saccharides?|[A-Za-z]*oses?)\b",
    flags=re.IGNORECASE,
)

CONNECTOR_RE = r"(?:of|with|in|on)(?:\s+(?:a|an|the|this|these|their))?"
LEADING_DETERMINER_RE = re.compile(
    r"^(?:a|an|the|this|these|their)\s+",
    flags=re.IGNORECASE,
)


def normalize_spaces(text: str) -> str:
    return re.sub(r"\s+", " ", str(text).strip())


def clean_bert_output_start(text: str) -> str:
    text = normalize_spaces(text)
    return LEADING_DETERMINER_RE.sub("", text, count=1)


def cell_text(value) -> str:
    if value is None:
        return ""
    return normalize_spaces(value)


def last_word(text: str) -> str:
    words = normalize_spaces(text).split()
    return words[-1].lower() if words else ""


def has_structure_word(text: str) -> bool:
    words = [w.lower() for w in normalize_spaces(text).split()]
    return any(w in STRUCTURE_WORDS for w in words)


def has_allowed_structure_term(text: str) -> bool:
    return bool(ALLOWED_STRUCTURE_TERM_RE.search(normalize_spaces(text)))


def valid_prefix_for_extended_preposition(prefix: str) -> bool:
    lw = last_word(prefix)
    return (
        lw in {"derivative", "derivatives"}
        or lw.endswith("tion")
        or lw.endswith("tions")
    )


def structure_case_allowed(
    first_np: str, second_np: str, merged_base_np: str, intermediate: str
) -> bool:
    if intermediate in WITH_ALLOWED:
        return True

    if has_structure_word(first_np) or has_structure_word(second_np):
        return has_allowed_structure_term(merged_base_np)

    return True


def previous_word_before(text: str, pos: int) -> str:
    before = text[:pos].rstrip()
    match = re.search(r"([A-Za-z0-9-]+)\W*$", before)
    return match.group(1).lower() if match else ""


def is_allowed_full_of_merge(text: str, connector_start: int) -> bool:
    word = previous_word_before(text, connector_start)
    return (
        word in {"derivative", "derivatives"}
        or word.endswith("tion")
        or word.endswith("tions")
    )


def text_after_of(text: str, connector_start: int) -> str:
    return normalize_spaces(text[connector_start + len("of") :])


def make_of_merge_values(merged_base_np: str, connector_start: int):
    merged_base_np = normalize_spaces(merged_base_np)
    if is_allowed_full_of_merge(merged_base_np, connector_start):
        return merged_base_np, merged_base_np
    return None


def has_comma_before_connector(text: str, connector_start: int) -> bool:
    return text[:connector_start].rstrip().endswith(",")


def find_merge_values(first_np: str, second_np: str, merged_base_np: str):
    if not first_np or not second_np or not merged_base_np:
        return None

    first_np = normalize_spaces(first_np)
    second_np = normalize_spaces(second_np)
    merged_base_np = normalize_spaces(merged_base_np)

    main_pattern = re.compile(
        rf"(?<!\S)({re.escape(first_np)})\s+"
        rf"({CONNECTOR_RE})\s+"
        rf"({re.escape(second_np)})(?!\S)",
        flags=re.IGNORECASE,
    )

    match = main_pattern.search(merged_base_np)
    if not match:
        return None

    intermediate = normalize_spaces(match.group(2)).lower()
    connector_start = match.start(2)
    connector_word = intermediate.split()[0]

    if connector_word in {"with", "in", "on"} and has_comma_before_connector(
        merged_base_np, connector_start
    ):
        return None

    if not structure_case_allowed(first_np, second_np, merged_base_np, intermediate):
        return None

    if intermediate in OF_ALLOWED:
        return make_of_merge_values(merged_base_np, connector_start)

    if intermediate in WITH_ALLOWED:
        direct_span = normalize_spaces(match.group(0))

        extended_pattern = re.compile(
            rf"(?<!\S)(.+?)\s+of\s+({re.escape(first_np)})\s+"
            rf"(with(?:\s+(?:a|an|the))?)\s+"
            rf"({re.escape(second_np)})(?!\S)",
            flags=re.IGNORECASE,
        )

        ext_match = extended_pattern.search(merged_base_np)
        if ext_match:
            prefix = normalize_spaces(ext_match.group(1))
            if valid_prefix_for_extended_preposition(prefix):
                value = normalize_spaces(ext_match.group(0))
                return value, value
            return direct_span, merged_base_np

        return direct_span, direct_span

    if intermediate in IN_ON_ALLOWED:
        if last_word(first_np) in STRUCTURE_WORDS:
            return normalize_spaces(merged_base_np), normalize_spaces(merged_base_np)

        extended_pattern = re.compile(
            rf"(?<!\S)(.+?)\s+of\s+({re.escape(first_np)})\s+"
            rf"((?:in|on)(?:\s+(?:a|an|the))?)\s+"
            rf"({re.escape(second_np)})(?!\S)",
            flags=re.IGNORECASE,
        )

        ext_match = extended_pattern.search(merged_base_np)
        if not ext_match:
            return None

        prefix = normalize_spaces(ext_match.group(1))
        if valid_prefix_for_extended_preposition(prefix):
            value = normalize_spaces(merged_base_np)
            return value, value

    return None


def find_merge_values_any_order(np1: str, np2: str, merged_base_np: str):
    values = find_merge_values(np1, np2, merged_base_np)
    if values:
        return values
    return find_merge_values(np2, np1, merged_base_np)


def required_headers(headers):
    return [
        PMID_COL,
        SENT_INDEX_COL,
        SENT_TEXT_COL,
        BERT_OUTPUT_COL,
        START_COL,
        END_COL,
        BERT_BASE_COL,
        MERGED_NP_COL,
    ]


def get_header_map(ws):
    return {
        str(cell.value).strip(): cell.column
        for cell in ws[1]
        if cell.value is not None
    }


def ensure_merge_column(ws, headers):
    if MERGE_FLAG_COL in headers:
        return headers

    merge_col_idx = ws.max_column + 1
    ws.cell(row=1, column=merge_col_idx).value = MERGE_FLAG_COL

    if merge_col_idx > 1:
        source = ws.cell(row=1, column=merge_col_idx - 1)
        target = ws.cell(row=1, column=merge_col_idx)
        if source.has_style:
            target._style = copy(source._style)
        target.font = copy(source.font)
        target.fill = copy(source.fill)
        target.border = copy(source.border)
        target.alignment = copy(source.alignment)
        target.number_format = source.number_format
        target.protection = copy(source.protection)

    return get_header_map(ws)


def find_offset(sent_text: str, value: str):
    sent_text = "" if sent_text is None else str(sent_text)
    value = normalize_spaces(value)

    if not sent_text or not value:
        return None

    exact_start = sent_text.find(value)
    if exact_start >= 0:
        return exact_start, exact_start + len(value)

    match = re.search(re.escape(value), sent_text, flags=re.IGNORECASE)
    if match:
        return match.start(), match.end()

    return None


def component_match_pattern(term: str):
    term = normalize_spaces(term)
    if not term:
        return None

    pieces = []
    tokens = term.split()
    for index, token in enumerate(tokens):
        piece = re.escape(token)
        if (
            index == len(tokens) - 1
            and re.search(r"[A-Za-z]$", token)
            and not token.lower().endswith("s")
        ):
            piece += "s?"
        pieces.append(piece)

    return re.compile(
        r"(?<![A-Za-z0-9])" + r"\s+".join(pieces) + r"(?![A-Za-z0-9])",
        flags=re.IGNORECASE,
    )


def find_component_span(text: str, term: str):
    text = normalize_spaces(text)
    term = normalize_spaces(term)
    if not text or not term:
        return None

    pattern = component_match_pattern(term)
    if pattern:
        match = pattern.search(text)
        if match:
            return match.start(), match.end()

    exact_start = text.find(term)
    if exact_start >= 0:
        return exact_start, exact_start + len(term)

    match = re.search(re.escape(term), text, flags=re.IGNORECASE)
    if match:
        return match.start(), match.end()

    return None


def merged_output_from_component_rows(ws, rows, headers, merged_np: str):
    spans = []
    merged_np = normalize_spaces(merged_np)

    for row_idx in rows:
        for col_name in (BERT_OUTPUT_COL, BERT_BASE_COL):
            term = cell_text(ws.cell(row=row_idx, column=headers[col_name]).value)
            span = find_component_span(merged_np, term)
            if span:
                spans.append(span)
                break

    if len(spans) < 2:
        return None

    start = min(span[0] for span in spans)
    end = max(span[1] for span in spans)
    if start >= end:
        return None

    return normalize_spaces(merged_np[start:end])


def group_fallback_offsets(ws, rows, headers):
    starts = []
    ends = []

    for row_idx in rows:
        start = ws.cell(row=row_idx, column=headers[START_COL]).value
        end = ws.cell(row=row_idx, column=headers[END_COL]).value
        try:
            starts.append(int(start))
            ends.append(int(end))
        except (TypeError, ValueError):
            continue

    if starts and ends:
        return min(starts), max(ends)

    return None


def copy_row_to_inserted_row(ws, source_row: int, target_row: int):
    for col_idx in range(1, ws.max_column + 1):
        source = ws.cell(row=source_row, column=col_idx)
        target = ws.cell(row=target_row, column=col_idx)
        target.value = source.value

        if source.has_style:
            target._style = copy(source._style)
        target.font = copy(source.font)
        target.fill = copy(source.fill)
        target.border = copy(source.border)
        target.alignment = copy(source.alignment)
        target.number_format = source.number_format
        target.protection = copy(source.protection)

    source_dim = ws.row_dimensions[source_row]
    target_dim = ws.row_dimensions[target_row]
    target_dim.height = source_dim.height
    target_dim.hidden = source_dim.hidden


def process_workbook():
    if not INPUT_FILE.exists():
        raise FileNotFoundError(f"Input file not found: {INPUT_FILE}")

    wb = load_workbook(INPUT_FILE)
    ws = wb[SHEET_NAME] if SHEET_NAME else wb.active

    headers = ensure_merge_column(ws, get_header_map(ws))

    for col in required_headers(headers):
        if col not in headers:
            raise ValueError(f"Column '{col}' not found in workbook.")

    merge_col_idx = headers[MERGE_FLAG_COL]
    for row_idx in range(2, ws.max_row + 1):
        ws.cell(row=row_idx, column=merge_col_idx).value = ""

    groups = {}
    for row_idx in range(2, ws.max_row + 1):
        merged_np = cell_text(ws.cell(row=row_idx, column=headers[MERGED_NP_COL]).value)
        bert_base_np = cell_text(ws.cell(row=row_idx, column=headers[BERT_BASE_COL]).value)

        if not merged_np or not bert_base_np:
            continue

        key = (
            ws.cell(row=row_idx, column=headers[PMID_COL]).value,
            ws.cell(row=row_idx, column=headers[SENT_INDEX_COL]).value,
            ws.cell(row=row_idx, column=headers[SENT_TEXT_COL]).value,
            merged_np.lower(),
        )
        groups.setdefault(key, []).append(row_idx)

    initial_data_rows = ws.max_row - 1
    rows_to_insert = []
    merged_group_count = 0

    for key, rows in groups.items():
        if len(rows) < 2:
            continue

        merged_np = cell_text(ws.cell(row=rows[0], column=headers[MERGED_NP_COL]).value)
        merge_values = None

        for i, row1 in enumerate(rows):
            np1 = cell_text(ws.cell(row=row1, column=headers[BERT_BASE_COL]).value)
            for row2 in rows[i + 1 :]:
                np2 = cell_text(ws.cell(row=row2, column=headers[BERT_BASE_COL]).value)
                merge_values = find_merge_values_any_order(np1, np2, merged_np)
                if merge_values:
                    break
            if merge_values:
                break

        if not merge_values:
            continue

        bert_output_value, _bert_base_value = merge_values
        merged_output_value = merged_output_from_component_rows(
            ws, rows, headers, merged_np
        )
        if merged_output_value is None:
            merged_output_value = clean_bert_output_start(bert_output_value)
        merged_base_value = merged_np
        template_row = min(rows)
        insert_after_row = max(rows)
        sent_text = ws.cell(row=template_row, column=headers[SENT_TEXT_COL]).value
        offsets = find_offset(sent_text, merged_output_value)
        if offsets is None:
            offsets = group_fallback_offsets(ws, rows, headers)

        rows_to_insert.append(
            (
                insert_after_row,
                template_row,
                merged_output_value,
                merged_base_value,
                offsets,
            )
        )
        merged_group_count += 1

    for (
        insert_after_row,
        template_row,
        merged_output_value,
        merged_base_value,
        offsets,
    ) in sorted(
        rows_to_insert, key=lambda item: item[0], reverse=True
    ):
        target_row = insert_after_row + 1
        ws.insert_rows(target_row, 1)
        copy_row_to_inserted_row(ws, template_row, target_row)

        ws.cell(row=target_row, column=headers[BERT_OUTPUT_COL]).value = merged_output_value
        ws.cell(row=target_row, column=headers[BERT_BASE_COL]).value = merged_base_value
        ws.cell(row=target_row, column=headers[MERGED_NP_COL]).value = merged_base_value
        ws.cell(row=target_row, column=headers[MERGE_FLAG_COL]).value = "Yes"

        if offsets is not None:
            start, end = offsets
            ws.cell(row=target_row, column=headers[START_COL]).value = start
            ws.cell(row=target_row, column=headers[END_COL]).value = end

    OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT_FILE)

    print(f"Input rows: {initial_data_rows}")
    print(f"Merged groups found: {merged_group_count}")
    print(f"Rows added: {len(rows_to_insert)}")
    print(f"Output rows: {ws.max_row - 1}")
    print(f"Saved output file to: {OUTPUT_FILE}")


if __name__ == "__main__":
    process_workbook()
