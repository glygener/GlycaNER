#!/usr/bin/env python3
"""
Split colon/slash BERT_Output and BERT_Base_NPs rows after BaseNP creation.

Rules:
1. Split one-token slash values when both BERT_Output and BERT_Base_NPs
   contain matching slash-separated parts.
   Example: GD1a/GT1b -> GD1a and GT1b.
2. Split colon values when BERT_Output contains colon-separated parts. Spaces
   around the colon are allowed. If BERT_Base_NPs does not have a matching part
   for one side, use that BERT_Output part as the BERT_Base_NPs value.
   Example: glycolipid: glycosyltransferase -> glycolipid and glycosyltransferase.
3. Add/update a ColonSlash column and mark "Yes" on the split output rows.
4. Update start_offset/end_offset from the original BERT_Output span.

Offsets are treated as 0-based, with end_offset exclusive.
"""

from __future__ import annotations
from glycaner_paths import STAGE3_INPUT, resource, work  # noqa: E402

import re
from copy import copy
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell


INFILE = Path(
    work("withBaseNPs_full.xlsx")
)
OUTFILE = Path(
    work("colon_slash_split.xlsx")
)
SHEET_NAME = None  # Use None to process every sheet with the required columns.

REQUIRED_COLUMNS = ["sent_text", "BERT_Output", "BERT_Base_NPs"]
TERM_COL = "BERT_Output"
BASE_COL = "BERT_Base_NPs"
SENT_TEXT_COL = "sent_text"
START_COL = "start_offset"
END_COL = "end_offset"
COLONSLASH_COL = "ColonSlash"


@dataclass
class Part:
    text: str
    rel_start: int
    rel_end: int


@dataclass
class SplitPlan:
    term_parts: list[Part]
    base_parts: list[Part]


def normalize_header(value: Any) -> str:
    return str(value).strip().lower() if value is not None else ""


def is_empty_cell(value: Any) -> bool:
    return value is None or str(value).strip() == ""


def clean_cell_text(value: Any) -> str:
    if is_empty_cell(value):
        return ""
    return str(value).strip()


def copy_cell_format(source: Cell, target: Cell) -> None:
    if source.has_style:
        target.font = copy(source.font)
        target.fill = copy(source.fill)
        target.border = copy(source.border)
        target.alignment = copy(source.alignment)
        target.protection = copy(source.protection)
        target.number_format = source.number_format
    if source.hyperlink:
        target._hyperlink = copy(source.hyperlink)
    if source.comment:
        target.comment = copy(source.comment)


def get_headers(ws) -> dict[str, int]:
    return {
        normalize_header(ws.cell(row=1, column=col).value): col
        for col in range(1, ws.max_column + 1)
    }


def ensure_colon_slash_column(ws, headers: dict[str, int]) -> dict[str, int]:
    key = normalize_header(COLONSLASH_COL)
    if key in headers:
        return headers

    col = ws.max_column + 1
    ws.cell(row=1, column=col).value = COLONSLASH_COL
    if col > 1:
        copy_cell_format(ws.cell(row=1, column=col - 1), ws.cell(row=1, column=col))

    headers[key] = col
    return headers


def has_required_columns(headers: dict[str, int]) -> bool:
    return all(normalize_header(column) in headers for column in REQUIRED_COLUMNS)


def read_row_template(ws, row_number: int, max_col: int) -> list[dict[str, Any]]:
    template = []
    for col in range(1, max_col + 1):
        cell = ws.cell(row=row_number, column=col)
        template.append(
            {
                "value": cell.value,
                "font": copy(cell.font),
                "fill": copy(cell.fill),
                "border": copy(cell.border),
                "alignment": copy(cell.alignment),
                "protection": copy(cell.protection),
                "number_format": cell.number_format,
                "hyperlink": copy(cell.hyperlink) if cell.hyperlink else None,
                "comment": copy(cell.comment) if cell.comment else None,
                "has_style": cell.has_style,
            }
        )
    return template


def write_row_from_template(ws, row_number: int, template: list[dict[str, Any]]) -> None:
    for col, info in enumerate(template, start=1):
        cell = ws.cell(row=row_number, column=col)
        cell.value = info["value"]
        if info["has_style"]:
            cell.font = copy(info["font"])
            cell.fill = copy(info["fill"])
            cell.border = copy(info["border"])
            cell.alignment = copy(info["alignment"])
            cell.protection = copy(info["protection"])
            cell.number_format = info["number_format"]
        if info["hyperlink"]:
            cell._hyperlink = copy(info["hyperlink"])
        if info["comment"]:
            cell.comment = copy(info["comment"])


def clone_template(template: list[dict[str, Any]]) -> list[dict[str, Any]]:
    return [{key: copy(value) for key, value in cell_info.items()} for cell_info in template]


def as_int(value: Any) -> int | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    try:
        return int(str(value).strip())
    except (TypeError, ValueError):
        return None


def find_span(
    sent_text: str,
    bert_output: str,
    start_offset: Any,
    end_offset: Any,
) -> tuple[int | None, int | None]:
    start = as_int(start_offset)
    end = as_int(end_offset)

    if start is not None and end is not None and 0 <= start <= end <= len(sent_text):
        if sent_text[start:end] == bert_output:
            return start, end
        if end - start == len(bert_output):
            return start, end

    found = sent_text.find(bert_output)
    if found != -1:
        return found, found + len(bert_output)

    lower_found = sent_text.casefold().find(bert_output.casefold())
    if lower_found != -1:
        return lower_found, lower_found + len(bert_output)

    return start, end


def split_with_spans(value: str, delimiter: str) -> list[Part]:
    parts: list[Part] = []
    part_start = 0

    for idx, char in enumerate(value):
        if char != delimiter:
            continue

        part = make_part(value[part_start:idx], part_start)
        if part is None:
            return []
        parts.append(part)
        part_start = idx + 1

    part = make_part(value[part_start:], part_start)
    if part is None:
        return []
    parts.append(part)
    return parts


def make_part(raw_text: str, raw_start: int) -> Part | None:
    left_trimmed = raw_text.lstrip()
    leading_removed = len(raw_text) - len(left_trimmed)
    text = left_trimmed.rstrip()
    if not text:
        return None

    rel_start = raw_start + leading_removed
    rel_end = rel_start + len(text)
    return Part(text=text, rel_start=rel_start, rel_end=rel_end)


def one_token_slash_value(value: str) -> bool:
    return "/" in value and not re.search(r"\s", value)


def normalized_contains(container: str, part: str) -> bool:
    container_norm = re.sub(r"\s+", " ", container.casefold()).strip()
    part_norm = re.sub(r"\s+", " ", part.casefold()).strip()
    return bool(container_norm and part_norm and part_norm in container_norm)


def infer_base_parts_from_term_parts(term_parts: list[Part], base_np: str) -> list[Part]:
    base_np = base_np.strip()
    base_part = make_part(base_np, 0) if base_np else None
    base_parts: list[Part] = []

    for term_part in term_parts:
        if base_part is not None and normalized_contains(base_part.text, term_part.text):
            base_parts.append(base_part)
        else:
            base_parts.append(
                Part(
                    text=term_part.text,
                    rel_start=term_part.rel_start,
                    rel_end=term_part.rel_end,
                )
            )

    return base_parts


def build_split_plan(bert_output: str, base_np: str) -> SplitPlan | None:
    if ":" in bert_output:
        term_parts = split_with_spans(bert_output, ":")
        if ":" in base_np:
            base_parts = split_with_spans(base_np, ":")
        else:
            base_parts = infer_base_parts_from_term_parts(term_parts, base_np)
    elif one_token_slash_value(bert_output) and one_token_slash_value(base_np):
        term_parts = split_with_spans(bert_output, "/")
        base_parts = split_with_spans(base_np, "/")
    else:
        return None

    if len(term_parts) < 2 or len(term_parts) != len(base_parts):
        return None

    return SplitPlan(term_parts=term_parts, base_parts=base_parts)


def process_sheet(ws) -> dict[str, int]:
    headers = get_headers(ws)
    if not has_required_columns(headers):
        return {"original_rows": ws.max_row - 1, "output_rows": ws.max_row - 1, "split_rows": 0, "added_rows": 0}

    headers = ensure_colon_slash_column(ws, headers)
    max_col = ws.max_column

    term_col = headers[normalize_header(TERM_COL)]
    base_col = headers[normalize_header(BASE_COL)]
    sent_col = headers[normalize_header(SENT_TEXT_COL)]
    start_col = headers.get(normalize_header(START_COL))
    end_col = headers.get(normalize_header(END_COL))
    colon_slash_col = headers[normalize_header(COLONSLASH_COL)]

    original_rows = ws.max_row - 1
    output_templates: list[list[dict[str, Any]]] = []
    split_rows = 0
    added_rows = 0

    for row_number in range(2, ws.max_row + 1):
        template = read_row_template(ws, row_number, max_col)

        bert_output = clean_cell_text(template[term_col - 1]["value"])
        base_np = clean_cell_text(template[base_col - 1]["value"])
        sent_text = clean_cell_text(template[sent_col - 1]["value"])

        plan = build_split_plan(bert_output, base_np)
        if plan is None:
            template[colon_slash_col - 1]["value"] = None
            output_templates.append(template)
            continue

        start_offset = template[start_col - 1]["value"] if start_col else None
        end_offset = template[end_col - 1]["value"] if end_col else None
        base_start, _ = find_span(sent_text, bert_output, start_offset, end_offset)

        split_rows += 1
        added_rows += len(plan.term_parts) - 1

        for term_part, base_part in zip(plan.term_parts, plan.base_parts):
            new_template = clone_template(template)
            new_template[term_col - 1]["value"] = term_part.text
            new_template[base_col - 1]["value"] = base_part.text
            new_template[colon_slash_col - 1]["value"] = "Yes"

            if base_start is not None and start_col and end_col:
                new_template[start_col - 1]["value"] = base_start + term_part.rel_start
                new_template[end_col - 1]["value"] = base_start + term_part.rel_end

            output_templates.append(new_template)

    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)

    for row_index, template in enumerate(output_templates, start=2):
        write_row_from_template(ws, row_index, template)

    if ws.auto_filter and ws.auto_filter.ref:
        ws.auto_filter.ref = ws.dimensions

    return {
        "original_rows": original_rows,
        "output_rows": len(output_templates),
        "split_rows": split_rows,
        "added_rows": added_rows,
    }


def process_workbook(
    input_path: Path,
    output_path: Path,
    sheet_name: str | None = None,
) -> dict[str, int]:
    wb = load_workbook(input_path)
    sheet_names = [sheet_name] if sheet_name else wb.sheetnames

    totals = {"original_rows": 0, "output_rows": 0, "split_rows": 0, "added_rows": 0}
    for name in sheet_names:
        summary = process_sheet(wb[name])
        for key, value in summary.items():
            totals[key] += value

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return totals


def main() -> None:
    input_path = INFILE.expanduser().resolve()
    output_path = OUTFILE.expanduser().resolve()

    summary = process_workbook(input_path, output_path, SHEET_NAME)

    print(f"Saved: {output_path}")
    print(f"Original data rows: {summary['original_rows']}")
    print(f"Output data rows: {summary['output_rows']}")
    print(f"Rows split by colon/slash: {summary['split_rows']}")
    print(f"Rows added by splitting: {summary['added_rows']}")


if __name__ == "__main__":
    main()
